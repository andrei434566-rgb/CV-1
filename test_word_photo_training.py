"""Number pairing, unabridged Word targets, and immutable portable exports."""

import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from PIL import Image

from app.infrastructure.word_photo_training import (
    export_word_photo_pairs, inspect_word_photo_pairs, read_numbered_descriptions,
)


class WordPhotoTrainingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.photos = self.root / "Фото"
        self.photos.mkdir()
        self.document = self.root / "Описания.docx"

    def photo(self, name):
        path = self.photos / name
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (20, 60), "gray").save(path)
        return path

    def word(self, *paragraphs):
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(self.document)

    def test_full_multiline_descriptions_and_excel_table_keep_document_order(self):
        first = "Песчаник серый. " * 1000
        document = Document()
        document.add_paragraph("Описания керна для обучения Qwen")
        document.add_paragraph("Фото 0001")
        document.add_paragraph(first)
        document.add_paragraph("")
        document.add_paragraph("Контакты чёткие.\nТонкая косая слоистость.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Фото № 0002.jpg"
        table.cell(0, 1).text = "Аргиллит тёмно-серый.\nБиотурбация слабая."
        document.add_paragraph("Фото 0003: Алевролит серый.")
        document.save(self.document)
        values = read_numbered_descriptions(self.document)
        self.assertEqual([number for number, _ in values], ["0001", "0002", "0003"])
        self.assertEqual(values[0][1], first + "\n\nКонтакты чёткие.\nТонкая косая слоистость.")
        self.assertEqual(values[1][1], "Аргиллит тёмно-серый.\nБиотурбация слабая.")
        self.assertEqual(values[2][1], "Алевролит серый.")

    def test_missing_duplicate_empty_and_corrupt_pairs_are_excluded(self):
        self.word("Фото 0001", "Первое.", "Фото 1", "Повторное.", "Фото 0002", "Фото 0003", "Нет фото.",
                  "Фото 0004", "Два файла.", "Фото 0005", "Повреждённое фото.", "Фото 0006", "Правильная пара.")
        for name in ("0001.jpg", "0002.jpg", "0004.jpg", "nested/4.png", "0006.jpg", "0009.jpg", "Керн.jpg"):
            self.photo(name)
        (self.photos / "0005.png").write_bytes(b"broken photo")
        inspected = inspect_word_photo_pairs(self.photos, self.document)
        self.assertEqual([pair.photo_id for pair in inspected.ready_pairs], ["0006"])
        messages = "\n".join(pair.message for pair in inspected.pairs)
        for expected in ("повторяется", "Пустое описание", "не найдена", "Несколько фотографий", "повреждено", "Нет описания", "Имя фото"):
            self.assertIn(expected, messages)

    def test_numeric_matching_is_exact_and_accepts_leading_zeros(self):
        self.word("Фото 0001", "Описание один.", "Фото 0010", "Описание десять.")
        self.photo("1.JPG")
        self.photo("00010.png")
        inspected = inspect_word_photo_pairs(self.photos, self.document)
        self.assertEqual([p.images[0].name for p in inspected.ready_pairs], ["1.JPG", "00010.png"])

    def test_export_requires_review_copies_originals_and_preserves_full_target(self):
        text = "Серая порода. " * 2000
        self.word("Фото 0001", text)
        original = self.photo("0001.jpg")
        inspected = inspect_word_photo_pairs(self.photos, self.document)
        with self.assertRaisesRegex(ValueError, "Подтвердите"):
            export_word_photo_pairs(inspected, self.root / "unapproved")
        result = export_word_photo_pairs(inspected, self.root / "portable", approved=True)
        row = json.loads(result["manifest"].read_text(encoding="utf-8"))
        self.assertEqual(row["conversations"][1]["value"], text.strip())
        self.assertEqual((result["manifest"].parent / row["image"]).read_bytes(), original.read_bytes())
        self.assertEqual((result["manifest"].parent / "descriptions.docx").read_bytes(), self.document.read_bytes())
        self.assertTrue(row["metadata"]["training_approved"])
        self.assertTrue(row["metadata"]["geology_approved"])
        self.assertNotIn("facies", row["metadata"])

    def test_changed_word_or_photo_cannot_be_silently_approved(self):
        self.word("Фото 0001", "Исходное.")
        image = self.photo("0001.jpg")
        inspected = inspect_word_photo_pairs(self.photos, self.document)
        Image.new("RGB", (20, 60), "red").save(image)
        with self.assertRaisesRegex(ValueError, "Фото 0001 изменилось"):
            export_word_photo_pairs(inspected, self.root / "changed_photo", approved=True)
        self.assertFalse((self.root / "changed_photo/word_photo_approved.jsonl").exists())
        self.word("Фото 0001", "Изменено.")
        with self.assertRaisesRegex(ValueError, "Word изменился"):
            export_word_photo_pairs(inspected, self.root / "changed_word", approved=True)

    def test_no_headers_fails_with_format_instructions(self):
        self.word("Обычная теория без номеров фото.")
        with self.assertRaisesRegex(ValueError, "Фото 0001"):
            inspect_word_photo_pairs(self.photos, self.document)


if __name__ == "__main__":
    unittest.main()
