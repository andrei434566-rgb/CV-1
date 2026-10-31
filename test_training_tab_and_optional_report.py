"""User workflows: optional report fields and Word-photo training from a tab."""

import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PIL import Image
from openpyxl import load_workbook
from PySide6.QtCore import QPointF, QProcess, QSettings
from PySide6.QtGui import QCloseEvent, QPixmap
from PySide6.QtWidgets import QApplication

from app.domain.models import FaciesDetection, PhotoRecord
from app.infrastructure.facies_postprocess import UNRECOGNIZED_FACIES
from app.ui.windows.main_window import MainWindow


_APP = QApplication.instance() or QApplication([])
_MODULE = "app.ui.windows.main_window"


class TrainingTabAndReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.settings = QSettings(str(self.root / "settings.ini"), QSettings.Format.IniFormat)
        settings_patch = patch(_MODULE + ".QSettings", return_value=self.settings)
        settings_patch.start()
        self.addCleanup(settings_patch.stop)
        self.window = MainWindow()
        self.addCleanup(self.window.close)

    def test_blank_project_exports_final_template_without_mandatory_inputs(self):
        path = self.root / "empty.xlsx"
        with (
            patch(_MODULE + ".QFileDialog.getSaveFileName", return_value=(str(path), "Книга Excel (*.xlsx)")),
            patch(_MODULE + ".QMessageBox.warning") as warning,
            patch(_MODULE + ".QMessageBox.information") as information,
            patch(_MODULE + ".QMessageBox.critical") as critical,
        ):
            self.window.export_core_description_report()
        warning.assert_not_called()
        information.assert_not_called()
        critical.assert_not_called()
        book = load_workbook(path)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.max_column, 22)
        self.assertEqual(sheet["V1"].value, "Краткое описание")
        self.assertIsNone(sheet["A4"].value)
        self.assertIsNone(sheet["S4"].value)
        self.assertEqual(sheet["R4"].value, '=IF(COUNT(M4:N4)=2,N4-M4,"")')

    def test_unresolved_facies_and_missing_columns_keep_depth_and_description(self):
        pixmap = QPixmap(30, 70)
        pixmap.fill()
        detection = FaciesDetection(UNRECOGNIZED_FACIES, .5,
            [QPointF(0, 0), QPointF(20, 0), QPointF(20, 60), QPointF(0, 60)],
            {"Название фации": UNRECOGNIZED_FACIES, "__vlm_description": "Алевролит серый."})
        detection.depth_from, detection.depth_to = 100, 101
        record = PhotoRecord("one", "one.png", pixmap, well_name="Р-31", detections=[detection])
        self.window._records = [record]
        path = self.root / "partial.xlsx"
        with (
            patch.object(self.window, "_show_stack"),
            patch.object(self.window.workspace, "render_scene_region") as render,
            patch(_MODULE + ".QFileDialog.getSaveFileName", return_value=(str(path), "Книга Excel (*.xlsx)")),
            patch(_MODULE + ".QMessageBox.warning") as warning,
            patch(_MODULE + ".QMessageBox.critical") as critical,
        ):
            self.window.export_core_description_report()
        render.assert_not_called()
        warning.assert_not_called()
        critical.assert_not_called()
        book = load_workbook(path)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual((sheet["M4"].value, sheet["N4"].value), (100, 101))
        self.assertIsNone(sheet["S4"].value)
        self.assertEqual(sheet["V4"].value, "Алевролит серый.")
        self.assertEqual(detection.label, UNRECOGNIZED_FACIES)
        self.assertEqual(detection.attributes["Название фации"], UNRECOGNIZED_FACIES)

    def test_unresolved_label_preserves_an_explicitly_filled_facies_name(self):
        pixmap = QPixmap(30, 70)
        detection = FaciesDetection(UNRECOGNIZED_FACIES, .5,
            [QPointF(0, 0), QPointF(20, 0), QPointF(20, 60)],
            {"Название фации": "Русловая", "Индекс фации": "F1"})
        self.window._records = [PhotoRecord("one", "one.png", pixmap, well_name="Р-31", detections=[detection])]
        with patch.object(self.window, "_show_stack"):
            wells = self.window._core_description_report_wells()
        self.assertEqual(wells[0]["layers"][0]["attributes"]["Название фации"], "Русловая")
        self.assertEqual(wells[0]["layers"][0]["attributes"]["Индекс фации"], "F1")

    def wait_for_job(self, page):
        deadline = time.monotonic() + 10
        while page.busy and time.monotonic() < deadline:
            _APP.processEvents()
            time.sleep(.005)
        self.assertFalse(page.busy, "Background import did not finish")

    def test_word_pairs_prepare_in_background_and_train_through_local_runner(self):
        self.window.show_training_center()
        center = self.window._training_center
        self.assertIs(self.window.sheets.currentWidget(), self.window._training_page)
        self.assertFalse(center.isWindow())
        self.window.show_training_center()
        self.assertIs(center, self.window._training_center)
        center.reject()
        self.assertFalse(center.isHidden())
        page = center.word_pairs
        photos = self.root / "Фото"
        photos.mkdir()
        Image.new("RGB", (30, 100), "gray").save(photos / "0001.jpg")
        document = Document()
        document.add_paragraph("Фото 0001")
        document.add_paragraph("Аргиллит серый.\nГоризонтальная слоистость.")
        source = self.root / "Описания.docx"
        document.save(source)
        page.photos.setText(str(photos))
        page.document.setText(str(source))
        page.check_pairs()
        self.wait_for_job(page)
        self.assertEqual(page.table.rowCount(), 1)
        self.assertIn("Аргиллит серый.", page.preview_text.toPlainText())
        self.assertFalse(page.export_button.isEnabled())
        page.approval.setChecked(True)
        with patch("app.ui.widgets.word_photo_training.user_data_root", return_value=self.root):
            page.export_pairs()
            self.wait_for_job(page)
        manifest = Path(str(self.settings.value("training/word_pairs_manifest")))
        self.assertTrue(manifest.is_file())
        self.assertTrue(page.train_button.isEnabled())
        self.assertEqual(json.loads(manifest.read_text(encoding="utf-8"))["metadata"]["photo_id"], "0001")
        base = self.root / "base"
        base.mkdir()
        (base / "config.json").write_text("{}")
        python = self.root / "python.exe"
        python.touch()
        process = Mock()
        process.state.return_value = QProcess.ProcessState.NotRunning
        process.waitForStarted.return_value = True
        center.process = process
        config = {"python": str(python), "base": str(base), "adapter": "", "theory_adapter": ""}
        with (
            patch.object(self.window, "_local_vlm_config", return_value=config),
            patch("app.ui.dialogs.training_center_dialog.user_data_root", return_value=self.root),
            patch.object(self.window._vlm_client, "stop") as stop,
        ):
            center._train("word_pairs")
        stop.assert_called_once()
        args = process.start.call_args.args[1]
        self.assertEqual(args[args.index("--manifest") + 1], str(manifest))
        self.assertEqual(args[args.index("--model") + 1], str(base))
        self.assertNotIn("--experimental-candidates", args)
        self.assertNotIn("--initial-adapter", args)
        self.assertFalse(page.train_button.isEnabled())
        page.set_training_busy(False)
        page.photos.setText(str(self.root / "другая папка"))
        self.assertFalse(page.train_button.isEnabled())
        self.assertIsNone(self.settings.value("training/word_pairs_manifest"))

    def test_main_window_does_not_destroy_running_training(self):
        self.window.show_training_center()
        center = self.window._training_center
        original = center.process
        process = Mock()
        process.state.return_value = QProcess.ProcessState.Running
        center.process = process
        try:
            event = QCloseEvent()
            self.window.closeEvent(event)
            self.assertFalse(event.isAccepted())
            with patch.object(self.window._vlm_client, "send") as send:
                self.window._drain_vlm_requests()
                self.window._send_vlm_chat("Опиши керн")
            send.assert_not_called()
        finally:
            center.process = original


if __name__ == "__main__":
    unittest.main()
