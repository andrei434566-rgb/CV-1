"""Incomplete reports stay exportable, with live Excel error highlighting."""
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from app.infrastructure.core_report_export import export_core_description_report


class ReportProblemHighlightTests(unittest.TestCase):
    def test_report_keeps_unknowns_and_regime_with_live_red_rules(self):
        well = {"name": "Р-31", "excel_report_inputs": {"gis_shift": 0, "penetration": 0, "core_recovered": 1},
                "layers": [{"label": "Не установлена по фото", "depth_from": 100, "depth_to": 101,
                "attributes": {"__vlm_description": "Серая порода.", "__vlm_hydrodynamic_regime": "Предположительно: спокойный"}}]}
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "report.xlsx"
            export_core_description_report(target, "Test", [well], "xlsx")
            book = load_workbook(target)
            try:
                sheet = book.active
                self.assertEqual(sheet.max_column, 22)
                self.assertEqual(sheet["G4"].value, 0)
                self.assertIn("Гидродинамический режим: Предположительно: спокойный", sheet["V4"].value)
                self.assertEqual(sheet["M4"].value, 100)
                rules = {str(key.sqref): value for key, value in sheet.conditional_formatting._cf_rules.items()}
                self.assertIn("S4", rules)
                self.assertIn("L4", rules)
                self.assertIn("G4", rules)
                for rule in rules["S4"]:
                    self.assertEqual(rule.type, "expression")
                    self.assertEqual(rule.dxf.fill.fgColor.rgb, "00FCE4E4")
                    self.assertIn("SEARCH", rule.formula[0])
                self.assertEqual(rules["G4"][0].formula, ['NOT(ISNUMBER(G4))'])
            finally:
                book.close()

    def test_invalid_drilling_passport_does_not_block_export_or_drop_layers(self):
        well = {"name": "Р-31", "excel_report_inputs": {"drilling_runs": [
                    {"number": "1", "top": 100, "base": 99}]},
                "layers": [{"label": "Русловая", "depth_from": 100, "depth_to": 101}]}
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "invalid.xlsx"
            export_core_description_report(target, "Test", [well], "xlsx")
            book = load_workbook(target)
            try:
                self.assertEqual(book.active["M4"].value, 100)
                self.assertIsNone(book.active["C4"].value)
                self.assertTrue(book.active.conditional_formatting)
            finally:
                book.close()

    def test_clean_blank_template_is_not_painted_red(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "empty.xlsx"
            export_core_description_report(target, "Test", [{"name": "Бланк", "layers": []}], "xlsx")
            book = load_workbook(target)
            try:
                self.assertEqual(len(book.active.conditional_formatting), 0)
            finally:
                book.close()

    def test_reversed_depths_remain_visible_for_correction(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "reversed.xlsx"
            export_core_description_report(target, "Test", [{"name": "Р-31", "layers": [
                {"label": "Русловая", "depth_from": 101, "depth_to": 100}]}], "xlsx")
            book = load_workbook(target)
            try:
                self.assertEqual((book.active["M4"].value, book.active["N4"].value), (101, 100))
                self.assertTrue(any("$N4<=$M4" in rule.formula[0]
                    for key, rules in book.active.conditional_formatting._cf_rules.items()
                    if str(key.sqref) == "M4" for rule in rules))
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
