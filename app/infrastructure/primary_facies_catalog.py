"""Portable, project-owned class vocabulary imported from a geologist's Excel."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path


SNAPSHOT_NAME = "primary_facies_catalog.json"
_HEADER_ALIASES = {
    "code": ("код фации", "фациальный код", "код", "facies code"),
    "index": ("индекс фации", "номер фации", "индекс", "№ фации", "facies index"),
    "name": ("название фации", "наименование фации", "имя фации", "название", "facies name"),
}


def _text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _header(value: object) -> str:
    return _text(value).casefold().replace("ё", "е").replace("_", " ")


def _header_columns(values: tuple) -> dict[str, int]:
    result: dict[str, int] = {}
    for column, value in enumerate(values):
        text = _header(value)
        for role, aliases in _HEADER_ALIASES.items():
            if role not in result and text in aliases:
                result[role] = column
    return result


def read_primary_facies_workbook(path: Path) -> dict:
    """Read a class dictionary; depth rows and photos are handled separately."""
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    classes: list[dict] = []
    seen: dict[str, dict] = {}
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            heading: tuple | None = None
            columns: dict[str, int] = {}
            for row_number, row in enumerate(rows, start=1):
                if heading is None:
                    candidate = _header_columns(row)
                    if ("code" in candidate or "name" in candidate) and len(candidate) >= 2:
                        heading, columns = row, candidate
                    elif row_number >= 30:
                        break
                    continue
                def cell(role: str) -> str:
                    column = columns.get(role)
                    return _text(row[column]) if column is not None and column < len(row) else ""
                code, index, name = cell("code"), cell("index"), cell("name")
                if not code and not name:
                    continue
                if code.casefold() in {"код фации", "facies code"}:
                    continue
                label = f"{code}@{index}" if code and index else code or name
                metadata = {
                    _text(title): _text(row[column])
                    for column, title in enumerate(heading)
                    if title is not None and column < len(row) and _text(row[column])
                }
                metadata.update({"Код фации": code, "Индекс фации": index, "Название фации": name})
                if label in seen:
                    existing = seen[label]
                    if existing["name"] != (name or code) and not index:
                        raise ValueError(f"Код «{label}» имеет разные названия без индекса ({sheet.title}!{row_number}).")
                    existing["source_variants"].append({"source": f"{sheet.title}!{row_number}", "metadata": metadata})
                    continue
                entry = {
                    "class_id": len(classes), "model_label": label,
                    "code": code, "index": index, "name": name or code,
                    "metadata": metadata, "source": f"{sheet.title}!{row_number}",
                    "source_variants": [],
                }
                classes.append(entry)
                seen[label] = entry
    finally:
        workbook.close()
    if not classes:
        raise ValueError("В Excel не найден справочник с заголовками «Код фации», «Индекс фации», «Название фации».")
    if len(classes) > 1000:
        raise ValueError("В справочнике больше 1000 строк классов. Проверьте выбранный лист и заголовки.")
    return {
        "schema": "excel-primary-facies-v1",
        "source_name": path.name,
        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "classes": classes,
    }


def load_catalog_snapshot(path: Path) -> dict | None:
    try:
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    classes = payload.get("classes") if isinstance(payload, dict) else None
    if not isinstance(payload, dict) or payload.get("schema") != "excel-primary-facies-v1" or not isinstance(classes, list) or not classes:
        return None
    if len({item.get("model_label") for item in classes if isinstance(item, dict)}) != len(classes):
        return None
    return payload


def save_catalog_snapshot(path: Path, catalog: dict) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def resolve_catalog_class(label: str, index: str | int | None, classes: list[dict]) -> dict | None:
    requested, requested_index = _text(label), _text(index)
    candidates = [item for item in classes if requested in (item["model_label"], item["code"], item["name"])]
    if requested_index:
        candidates = [item for item in candidates if item["index"] == requested_index]
    return candidates[0] if len(candidates) == 1 else None
