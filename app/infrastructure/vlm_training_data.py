"""Export only geologist-approved image descriptions and explicit theory Q&A."""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtGui import QImage, QPainter, QPainterPath, QPolygonF


def _description(detection) -> str:
    attrs = detection.attributes or {}
    keys = ("__vlm_description", "Краткое описание", "Литологическое описание (16 параметров)") if attrs.get("__geometry_only") == "1" else ("Краткое описание", "Литологическое описание (16 параметров)")
    return next((str(attrs.get(key) or "").strip() for key in keys if str(attrs.get(key) or "").strip()), "")


def _expert_description_fields(detection) -> dict[str, str]:
    return {
        str(key): str(value)
        for key, value in (detection.attributes or {}).items()
        if not str(key).startswith("__") or str(key).startswith("__vlm_")
    }


def expert_description_digest(detection) -> str:
    """Bind an explicit expert confirmation to this exact text and geometry.

    Displayed attributes and the VLM's interpretation fields participate in the
    digest; transient internal IDs and the confirmation stamps themselves do
    not. Any later text, palette, facies, polygon or depth edit invalidates it.
    """
    payload = {
        "label": str(detection.label), "attributes": _expert_description_fields(detection),
        "polygon": [[float(point.x()), float(point.y())] for point in detection.polygon],
        "depth_from": detection.depth_from, "depth_to": detection.depth_to,
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def expert_description_is_confirmed(detection) -> bool:
    attrs = detection.attributes or {}
    if not (
        attrs.get("__geometry_only") == "1"
        and detection.training_ready
        and attrs.get("__expert_description_confirmed") == "1"
        and attrs.get("__expert_description_sha256")
        and _description(detection)
    ):
        return False
    try:
        return attrs["__expert_description_sha256"] == expert_description_digest(detection)
    except (TypeError, ValueError):
        return False


def _pair_provenance(detection) -> tuple[str, str]:
    """Keep legacy direct-Excel targets and explicitly confirmed CV-1 targets."""
    attrs = detection.attributes or {}
    if attrs.get("__core_column_description") == "1":
        return "", "Описание относится ко всему столбику, а не к отдельной фации"
    if not _description(detection):
        return "", "Нет непустого описания породы"
    if len(detection.polygon) < 3:
        return "", "Нет корректного контура"
    if attrs.get("__geometry_only") == "1":
        if not expert_description_is_confirmed(detection):
            return "", "Описание CV-1 не подтверждено отдельно геологом или изменилось после подтверждения"
        return "expert_cv1_reviewed_mask", ""
    if str(attrs.get("Источник описания", "")).startswith("Excel:"):
        return "source_direct_excel_reviewed_mask", ""
    return "", "Нет исходного описания Excel или отдельно подтверждённого описания CV-1"


def reviewed_pair_counts(records) -> tuple[int, int]:
    reviewed = usable = 0
    for record in records:
        for detection in record.detections:
            if not detection.training_ready:
                continue
            reviewed += 1
            provenance, _ = _pair_provenance(detection)
            if provenance:
                usable += 1
    return reviewed, usable


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def export_reviewed_pairs(records, destination: Path, *, include_facies: bool = False) -> dict:
    """Save direct Excel or separately expert-confirmed CV-1 image/text pairs."""
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError(f"Учебная папка уже существует: {destination}")
    crops = destination / "crops"
    crops.mkdir(parents=True)
    rows = []
    skipped = []
    counts = {"source_direct_excel_reviewed_mask": 0, "expert_cv1_reviewed_mask": 0}
    for record in records:
        if record.pixmap.isNull():
            continue
        image = record.pixmap.toImage().convertToFormat(QImage.Format.Format_RGB32)
        for index, detection in enumerate(record.detections):
            if not detection.training_ready:
                continue
            attrs = detection.attributes or {}
            description = _description(detection)
            provenance, reason = _pair_provenance(detection)
            if not provenance:
                skipped.append({"photo": record.path, "layer": index, "reason": reason})
                continue
            polygon = QPolygonF(detection.polygon)
            rect = polygon.boundingRect().toAlignedRect().intersected(image.rect())
            if rect.width() < 8 or rect.height() < 8:
                skipped.append({"photo": record.path, "layer": index, "reason": "Слишком маленький контур"})
                continue
            crop = QImage(rect.size(), QImage.Format.Format_RGB32)
            crop.fill(Qt.GlobalColor.white)
            painter = QPainter(crop)
            path = QPainterPath()
            path.addPolygon(polygon.translated(-rect.x(), -rect.y()))
            painter.setClipPath(path)
            painter.drawImage(-rect.x(), -rect.y(), image)
            painter.end()
            name = f"pair_{len(rows)+1:06d}.jpg"
            if not crop.save(str(crops / name), "JPG", 94):
                raise OSError(f"Не удалось сохранить вырезку: {name}")
            question = "<image>Опиши видимые признаки породы на русском языке. Не выдумывай невидимые детали."
            target = description
            if include_facies and provenance == "expert_cv1_reviewed_mask":
                # Both the facies/contour and the text have independent,
                # current specialist approvals in _pair_provenance above.
                question = ("<image>Опиши видимые признаки породы и предложи фацию по-русски. "
                            "Отделяй наблюдения от интерпретации, не выдумывай невидимые признаки.")
                target += "\nФация: " + str(detection.label).strip()
                for field, caption in (("__vlm_facies_association", "Ассоциация фаций"),
                                       ("__vlm_depositional_environment", "Обстановка осадконакопления")):
                    value = str(attrs.get(field) or "").strip()
                    if value:
                        target += "\n" + caption + ": " + value
            rows.append({
                "image": f"crops/{name}",
                "conversations": [
                    {"from": "human", "value": question},
                    {"from": "gpt", "value": target},
                ],
                "metadata": {
                    "training_approved": True,
                    "geology_approved": True,
                    "record_type": "image_text",
                    "label_provenance": provenance,
                    "source_photo": str(record.path),
                    "source_well": record.well_name,
                    "source_row": attrs.get("Источник описания", ""),
                    "source_geometry": attrs.get("Источник геометрии", ""),
                    "source_geometry_id": attrs.get("__geometry_id", ""),
                    "source_model_class": attrs.get("__source_model_class", ""),
                    "source_model_class_id": attrs.get("__source_model_class_id", ""),
                    "source_model_sha256": attrs.get("__source_model_sha256", ""),
                    "facies_in_target": bool(include_facies and provenance == "expert_cv1_reviewed_mask"),
                    "expert_description_sha256": attrs.get("__expert_description_sha256", "") if provenance == "expert_cv1_reviewed_mask" else "",
                    "expert_reviewed_fields": _expert_description_fields(detection) if provenance == "expert_cv1_reviewed_mask" else {},
                    "facies": detection.label,
                    "depth_top": detection.depth_from,
                    "depth_base": detection.depth_to,
                },
            })
            counts[provenance] += 1
    if not rows:
        raise ValueError("Нет проверенных контуров с исходным описанием Excel или отдельно подтверждённым описанием CV-1. Проверьте слои и текст в карточках фаций.")
    manifest = destination / "photo_excel_approved.jsonl"
    _write_jsonl(manifest, rows)
    (destination / "export_report.json").write_text(json.dumps({
        "schema": "kern-vlm-reviewed-image-text-v2",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "approved_pairs": len(rows), "skipped": skipped,
        "pairs_by_provenance": counts,
        "warning": "Для CV-1 нужны отдельное подтверждение текста геологом и неизменные после него описание и контур. Автоматический ответ VLM не становится учебной целью сам. Контрольное качество оценивайте отдельно по независимым скважинам.",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"manifest": manifest, "count": len(rows), "skipped": len(skipped), "pairs_by_provenance": counts}


_QUESTION = {"вопрос", "question", "запрос"}
_ANSWER = {"ответ", "answer"}


def _qa_rows(path: Path) -> list[tuple[str, str, str]]:
    result = []
    if path.suffix.casefold() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [str(value or "").strip().casefold() for value in next(rows, ())]
            q = next((i for i, h in enumerate(header) if h in _QUESTION), None)
            a = next((i for i, h in enumerate(header) if h in _ANSWER), None)
            if q is None or a is None:
                continue
            for row_number, cells in enumerate(rows, start=2):
                question = str(cells[q] or "").strip() if q < len(cells) else ""
                answer = str(cells[a] or "").strip() if a < len(cells) else ""
                if question and answer:
                    result.append((question, answer, f"{sheet.title}!{row_number}"))
        book.close()
    elif path.suffix.casefold() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(2048)
            stream.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            reader = csv.DictReader(stream, dialect=dialect)
            for number, row in enumerate(reader, start=2):
                normalized = {str(k).strip().casefold(): str(v or "").strip() for k, v in row.items() if k}
                question = next((normalized[k] for k in _QUESTION if normalized.get(k)), "")
                answer = next((normalized[k] for k in _ANSWER if normalized.get(k)), "")
                if question and answer:
                    result.append((question, answer, f"строка {number}"))
    else:
        raise ValueError("Для отдельного обучения теории нужна таблица .xlsx или .csv с колонками «Вопрос» и «Ответ».")
    return result


def export_theory_qa(source: Path, destination: Path) -> dict:
    """Convert only user-authored explicit answers, never inferred geology."""
    source = Path(source)
    pairs = _qa_rows(source)
    if not pairs:
        raise ValueError("В таблице не найдены заполненные пары «Вопрос» — «Ответ».")
    records = [{
        "conversations": [
            {"from": "human", "value": question + " Отвечай по-русски; если данных недостаточно, сообщи об этом."},
            {"from": "gpt", "value": answer},
        ],
        "metadata": {
            "training_approved": True, "record_type": "text_only", "source_group": str(source),
            "source_row": address, "label_provenance": "user_supplied_theory_qa",
        },
    } for question, answer, address in pairs]
    _write_jsonl(destination, records)
    return {"manifest": destination, "count": len(records)}
