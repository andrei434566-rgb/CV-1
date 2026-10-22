"""Stage source-direct UKCS/IODP interval material for human review, never training.

No archive member is executed and no translated geological target is invented.
UKCS labels are row-registered class codes, not descriptions. IODP tables and
documents remain searchable even when a safe photo join is unavailable.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import re
import shutil
import stat
import zipfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path, PurePosixPath


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
DOCUMENT_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".md", ".docx", ".pdf"}
MAX_DOCUMENT_BYTES = 64 * 1024 * 1024
MAX_VECTOR_BYTES = 128 * 1024 * 1024
MAX_TABLE_ROWS = 200_000
MAX_ARCHIVE_ENTRIES = 50_000
UKCS_URL = "https://figshare.com/articles/dataset/UKCS_Q204_Q205_Subsurface_Data_products_for_Machine_Learning_Study/14265689"
IODP_URL = "https://publications.iodp.org/proceedings/362/362title.html"
IODP_METHODS_URL = "https://publications.iodp.org/proceedings/362/102/362_102.html"


class DatasetPreparationCancelled(RuntimeError):
    """Raised at a safe boundary; partial review files are not trainable."""


def _check(cancel) -> None:
    if cancel is not None and (cancel() if callable(cancel) else cancel.is_set()):
        raise DatasetPreparationCancelled("Подготовка источника отменена")


def _progress(callback, message: str) -> None:
    if callback:
        callback(message)


def _id(*values) -> str:
    return hashlib.sha256("\0".join(str(v) for v in values).encode("utf-8")).hexdigest()[:24]


def _jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text("".join(json.dumps(row, ensure_ascii=False, default=str) + "\n" for row in rows), encoding="utf-8")


@dataclass(frozen=True)
class _Entry:
    path: Path
    member: str = ""
    size: int = 0

    @property
    def name(self) -> str:
        return self.member or self.path.name

    @property
    def reference(self) -> str:
        return f"{self.path}!{self.member}" if self.member else str(self.path)

    @contextmanager
    def open(self):
        if self.member:
            with zipfile.ZipFile(self.path) as archive, archive.open(self.member) as stream:
                yield stream
        else:
            with self.path.open("rb") as stream:
                yield stream


def _safe_member(info: zipfile.ZipInfo) -> bool:
    name = info.filename.replace("\\", "/")
    parts = name.split("/")
    return bool(name and not name.startswith("/") and not re.match(r"^[A-Za-z]:", name)
                and all(p not in {"..", "."} for p in parts)
                and not any(":" in p for p in parts)
                and not stat.S_ISLNK(info.external_attr >> 16) and not (info.flag_bits & 1))


def _inventory(paths: list[Path], output_dir: Path, unmatched: list[dict], cancel) -> list[_Entry]:
    entries, seen = [], set()
    files = []
    for value in paths:
        _check(cancel)
        path = Path(value).resolve()
        if not path.exists():
            unmatched.append({"source_path": str(path), "reason": "source_missing"})
            continue
        if path == output_dir or output_dir in path.parents:
            raise ValueError("Выходная папка не может быть входным источником")
        if path.is_dir():
            files.extend(p for p in sorted(path.rglob("*")) if p.is_file() and not p.is_symlink()
                         and output_dir not in p.resolve().parents)
        else:
            files.append(path)
    for path in sorted(set(files)):
        _check(cancel)
        if path.suffix.casefold() == ".zip":
            try:
                with zipfile.ZipFile(path) as archive:
                    infos = archive.infolist()
                    if len(infos) > MAX_ARCHIVE_ENTRIES:
                        raise ValueError("archive_entry_limit")
                    for info in infos:
                        if info.is_dir():
                            continue
                        if not _safe_member(info):
                            unmatched.append({"source_path": str(path), "source_member": info.filename,
                                              "reason": "unsafe_archive_member"})
                            continue
                        entry = _Entry(path, info.filename, info.file_size)
                        if entry.reference in seen:
                            unmatched.append({"source_path": entry.reference, "reason": "duplicate_archive_member"})
                            continue
                        seen.add(entry.reference)
                        entries.append(entry)
            except (zipfile.BadZipFile, OSError, ValueError) as exc:
                unmatched.append({"source_path": str(path), "reason": "archive_unreadable", "detail": str(exc)})
        else:
            entry = _Entry(path, size=path.stat().st_size)
            if entry.reference not in seen:
                entries.append(entry)
                seen.add(entry.reference)
    return entries


def _npy_header(stream):
    import numpy as np
    version = np.lib.format.read_magic(stream)
    if version[0] == 1:
        return np.lib.format.read_array_header_1_0(stream, max_header_size=65536)
    if version[0] in {2, 3}:
        return np.lib.format.read_array_header_2_0(stream, max_header_size=65536)
    raise ValueError(f"Unsupported NPY version: {version}")


def _vector(entry: _Entry, kinds: str):
    import numpy as np
    with entry.open() as stream:
        shape, _, dtype = _npy_header(stream)
        if len(shape) != 1 or dtype.hasobject or dtype.kind not in kinds:
            raise ValueError("Expected a non-object one-dimensional source array")
        count = math.prod(shape)
        if count * dtype.itemsize > MAX_VECTOR_BYTES:
            raise ValueError("Source vector exceeds memory limit")
        raw = stream.read(count * dtype.itemsize + 1)
        if len(raw) != count * dtype.itemsize:
            raise ValueError("Truncated or surplus source vector bytes")
        return np.frombuffer(raw, dtype=dtype).copy()


def _base_record(source_id: str, url: str, entry: _Entry, address: str) -> dict:
    return {"record_id": _id(source_id, entry.reference, address), "source_id": source_id,
            "source_url": url, "source_path": str(entry.path), "source_member": entry.member,
            "source_row": address, "training_approved": False, "geology_approved": False,
            "review_status": "source_direct_unreviewed", "translation_status": "original_only"}


def _skip(stream, count: int, cancel) -> None:
    while count:
        _check(cancel)
        data = stream.read(min(count, 4 * 1024 * 1024))
        if not data:
            raise ValueError("Truncated source image array")
        count -= len(data)


def _ukcs(source_id, entries, output_dir, intervals, candidates, unmatched, progress, cancel):
    import numpy as np
    from PIL import Image
    by_location = {(str(e.path), e.member or e.path.name): e for e in entries}
    used = set()
    for labels_entry in [e for e in entries if e.name.endswith("_labels.npy")]:
        _check(cancel)
        prefix = labels_entry.name.removesuffix("_labels.npy")
        well = PurePosixPath(prefix).name
        _progress(progress, f"UKCS: исходные интервалы {well}")
        key = str(labels_entry.path)
        # Loose triplets have different paths; ZIP triplets share one archive.
        def sibling(suffix):
            name = prefix + suffix
            if labels_entry.member:
                return by_location.get((key, name))
            return next((e for e in entries if not e.member and e.path == labels_entry.path.with_name(PurePosixPath(name).name)), None)
        depth_entry, image_entry = sibling("_depth.npy"), sibling("_image.npy")
        used.add(labels_entry.reference)
        if not depth_entry or not image_entry:
            unmatched.append({"source_path": labels_entry.reference, "reason": "missing_depth_or_image_triplet", "well_id": well})
            continue
        used.update((depth_entry.reference, image_entry.reference))
        try:
            labels = _vector(labels_entry, "SU")
            depth = _vector(depth_entry, "fiu")
            with image_entry.open() as stream:
                shape, fortran, dtype = _npy_header(stream)
                image_header_bytes = stream.tell()
            if len(labels) != len(depth) or not len(labels) or not np.isfinite(depth).all():
                raise ValueError("Label/depth sizes differ, are empty or contain nonfinite depths")
            if fortran or len(shape) != 3 or shape[0] != len(depth) or shape[2] != 3 or dtype != np.dtype("uint8"):
                raise ValueError("Image rows are not exactly aligned to source label/depth rows")
            if shape[1] < 1 or shape[1] > 20_000:
                raise ValueError("Unsupported image width")
            if math.prod(shape) * dtype.itemsize + image_header_bytes != image_entry.size:
                raise ValueError("Source image size does not match its NPY header")
            changes = (labels[1:] != labels[:-1]) | (np.diff(depth) < 0) | (np.diff(depth) > 0.01)
            bounds = np.r_[0, np.flatnonzero(changes) + 1, len(labels)]
            vector_digest = hashlib.sha256()
            vector_digest.update(labels.tobytes())
            vector_digest.update(depth.tobytes())
            vector_sha256 = vector_digest.hexdigest()
            well_intervals = []
            for start, end in zip(bounds[:-1], bounds[1:]):
                start, end = int(start), int(end)
                raw_code = labels[start]
                code = raw_code.decode("utf-8", errors="strict") if isinstance(raw_code, bytes) else str(raw_code)
                record = _base_record(source_id, UKCS_URL, labels_entry, f"rows[{start}:{end}]")
                record.update(well_id=well, source_group_id=f"{source_id}:{well}",
                              source_depth_group_id=f"{source_id}:{well}:{start}:{end}",
                              source_label=code, original_description="", target_kind="source_class_code",
                              row_top=start, row_bottom=end, depth_top_m=float(depth[start]),
                              depth_bottom_m=float(depth[end - 1]), depth_reference="source_depth.npy",
                              depth_bounds_kind="first_and_last_observed_row_not_interval_edges",
                              image_source=image_entry.reference, image="", label_source=labels_entry.reference,
                              depth_source=depth_entry.reference, pairing_status="source_arrays_row_aligned",
                              source_vector_sha256=vector_sha256,
                              alignment_method="Equal row counts; runs split on class changes, decreasing depth or gaps >0.01 m")
                intervals.append(record)
                well_intervals.append(record)
            # Bounded preview per code; all original runs remain in intervals.jsonl.
            groups = defaultdict(list)
            for row in well_intervals:
                if row["row_bottom"] - row["row_top"] >= 8:
                    groups[row["source_label"]].append(row)
            selected = []
            for group in groups.values():
                selected.extend(sorted(group, key=lambda r: (-(r["row_bottom"] - r["row_top"]), r["row_top"]))[:2])
            windows = []
            for row in selected:
                height = min(600, row["row_bottom"] - row["row_top"])
                top = (row["row_top"] + row["row_bottom"] - height) // 2
                windows.append((top, top + height, row))
            with image_entry.open() as stream:
                _npy_header(stream)
                cursor, row_bytes = 0, shape[1] * 3
                for top, bottom, row in sorted(windows, key=lambda item: item[0]):
                    _check(cancel)
                    _skip(stream, (top - cursor) * row_bytes, cancel)
                    raw = stream.read((bottom - top) * row_bytes)
                    if len(raw) != (bottom - top) * row_bytes:
                        raise ValueError("Truncated source crop")
                    image = np.frombuffer(raw, dtype=np.uint8).reshape(bottom - top, shape[1], 3)
                    relative = f"images/{row['record_id']}.png"
                    Image.fromarray(image).save(output_dir / relative)
                    metadata = dict(row, crop_row_top=top, crop_row_bottom=bottom,
                                    crop_depth_top_m=float(depth[top]), crop_depth_bottom_m=float(depth[bottom - 1]),
                                    source_run_depth_top_m=row["depth_top_m"], source_run_depth_bottom_m=row["depth_bottom_m"],
                                    depth_top_m=float(depth[top]), depth_bottom_m=float(depth[bottom - 1]),
                                    image=relative, pixel_alignment_verified=True)
                    candidates.append({"image": relative, "source_label": row["source_label"],
                                       "original_description": "", "metadata": metadata,
                                       "conversations": [{"from": "human", "value": "<image>Назови исходный код литологического класса UKCS."},
                                                         {"from": "gpt", "value": row["source_label"]}]})
                    row["preview_image"] = relative
                    cursor = bottom
        except (ValueError, OSError, EOFError, UnicodeError, zipfile.BadZipFile) as exc:
            unmatched.append({"source_path": labels_entry.reference, "well_id": well,
                              "reason": "ukcs_alignment_or_preview_failed", "detail": str(exc)})
    return used


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold().replace("ё", "е")


def _number(value):
    try:
        if isinstance(value, bool):
            return None
        result = float(str(value).strip().replace(",", "."))
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def _normalized_ooxml(data: bytes) -> tuple[bytes, list[dict]]:
    """Repair two known IODP ZIP-part case errors in memory only.

    OOXML part names are case sensitive. Some upstream DESClogik workbooks
    store lowercase part names while relationships use the canonical spelling.
    No XML or cell value is rewritten, and conflicting case variants are rejected.
    """
    canonical = {"[content_types].xml": "[Content_Types].xml",
                 "xl/sharedstrings.xml": "xl/sharedStrings.xml"}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_ENTRIES or sum(info.file_size for info in infos) > MAX_DOCUMENT_BYTES:
            raise ValueError("workbook_uncompressed_size_or_entry_limit")
        seen, changes = set(), []
        for info in infos:
            if not _safe_member(info):
                raise ValueError("unsafe_workbook_part")
            folded = info.filename.casefold()
            if folded in seen:
                raise ValueError("conflicting_case_or_duplicate_workbook_part")
            seen.add(folded)
            target = canonical.get(folded, info.filename)
            if target != info.filename:
                changes.append({"from": info.filename, "to": target,
                                "operation": "in_memory_ooxml_part_case_normalization"})
        if not changes:
            return data, []
        target_data = io.BytesIO()
        with zipfile.ZipFile(target_data, "w", compression=zipfile.ZIP_DEFLATED) as normalized:
            for info in infos:
                target = canonical.get(info.filename.casefold(), info.filename)
                normalized.writestr(target, archive.read(info))
        return target_data.getvalue(), changes


def _table_rows(entry: _Entry, normalizations: list[dict] | None = None):
    """Yield sheet, source row, original header/cells, including unrecognized tables."""
    suffix = PurePosixPath(entry.name).suffix.casefold()
    with entry.open() as stream:
        data = stream.read(MAX_DOCUMENT_BYTES + 1)
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError("document_size_limit")
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        data, changes = _normalized_ooxml(data)
        if normalizations is not None:
            normalizations.extend(changes)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                header = None
                for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    if number > MAX_TABLE_ROWS:
                        raise ValueError("table_row_limit")
                    if not any(value is not None for value in cells):
                        continue
                    # Preserve leading title rows as searchable text. A header is
                    # selected only when it exposes named fields, never by numbers.
                    if header is None:
                        if sum(bool(_norm(v)) for v in cells) >= 2 and any(_field(v) for v in cells):
                            header = [str(v or f"column_{i+1}") for i, v in enumerate(cells)]
                        else:
                            yield sheet.title, number, {f"column_{i+1}": v for i, v in enumerate(cells) if v is not None}
                        continue
                    yield sheet.title, number, {f"{key} [{i+1}]": cells[i] if i < len(cells) else None for i, key in enumerate(header)}
        finally:
            workbook.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        header = next(reader, ())
        for number, cells in enumerate(reader, 2):
            if number > MAX_TABLE_ROWS:
                raise ValueError("table_row_limit")
            if not any(str(value).strip() for value in cells):
                continue
            yield "", number, {f"{key} [{i+1}]": cells[i] if i < len(cells) else "" for i, key in enumerate(header)}


def _field(header):
    text = re.sub(r" \[\d+\]$", "", _norm(header))
    if text in {"image filename", "image file", "photo filename", "photo", "image", "фото", "файл фото"}:
        return "image"
    if (text.startswith(("top depth", "depth top")) or text in {"кровля", "глубина от", "depth_from"}) and not re.search(r"\b(?:cm|mm)\b", text):
        return "top"
    if (text.startswith(("bottom depth", "depth bottom")) or text in {"подошва", "глубина до", "depth_to"}) and not re.search(r"\b(?:cm|mm)\b", text):
        return "bottom"
    if text in {"full lithology", "lithology", "principal lithology", "lithology principal name", "simplified lithology", "литология"}:
        return "label"
    if text in {"description", "lithologic description", "lithological description", "описание", "краткое описание",
                "general interval comment", "section summary description", "core summary description", "unit summary description"}:
        return "description"
    if text == "sample":
        return "sample id"
    if text == "exp":
        return "expedition"
    if text == "a/w":
        return "half"
    if text in {"site", "hole", "core", "type", "sect", "section", "half", "expedition", "sample id", "sample identifier", "well", "well id"}:
        return text
    return None


def _fields(raw):
    result = {}
    for key, value in raw.items():
        kind = _field(key)
        if kind and value is not None and str(value).strip():
            # Multiple depth scales are ambiguous; do not silently choose one.
            if kind in result and result[kind][1] != value:
                result[kind] = ("ambiguous", None)
            else:
                result[kind] = (key, value)
    return result


def _identity(fields):
    def val(key):
        return str(fields.get(key, ("", ""))[1] or "").strip().upper()
    def ordinal(value):
        return str(int(value)) if re.fullmatch(r"\d+", value) else value
    if val("sample id") or val("sample identifier"):
        sample = val("sample id") or val("sample identifier")
        match = re.fullmatch(r"(\d+)-(U\d+)([A-Z])-(\d+)([A-Z])-(\d+|CC)(?:-([AW]))?", sample)
        if match:
            expedition, site, hole, core, core_type, section, half = match.groups()
            return ("iodp", expedition, site, hole, ordinal(core), core_type, ordinal(section), half or "")
        return ("sample", sample)
    if val("site") and val("hole") and val("core") and val("type"):
        return ("iodp", val("expedition"), val("site"), val("hole"), ordinal(val("core")),
                val("type"), ordinal(val("sect") or val("section")), val("half"))
    return ()


def _iodp_source_fields(entry: _Entry, raw: dict, fields: dict) -> list[dict]:
    """Recognize published 362 LSIMG columns without guessing a depth scale.

    The generic parser deliberately rejects multiple depth scales. LSIMG has
    explicit CSF-A and CSF-B columns; Expedition 362 documents CSF-A as the
    standard sample/measurement scale. Preserve every original cell and record
    each source-specific selection. Unrecognized tables retain generic rules.
    """
    name = PurePosixPath(entry.name).name
    if not re.fullmatch(r"362-U\d+[A-Z]-LSIMG\.csv", name, re.I):
        return []
    by_header = {re.sub(r" \[\d+\]$", "", _norm(key)): (key, value) for key, value in raw.items()}
    selections = []
    for kind, header in (("top", "top depth csf-a (m)"), ("bottom", "bottom depth csf-a (m)")):
        field = by_header.get(header)
        if field is None or _number(field[1]) is None:
            return []
    for kind, header in (("top", "top depth csf-a (m)"), ("bottom", "bottom depth csf-a (m)")):
        fields[kind] = by_header[header]
        selections.append({"field": kind, "source_header": by_header[header][0],
                           "rule": "explicit_iodp362_lsimg_csf_a_column", "source_url": IODP_METHODS_URL})
    # The trimmed source image is preferred; no crop is generated by this module.
    for header in ("cropped image filename", "uncropped image filename"):
        field = by_header.get(header)
        if field and str(field[1] or "").strip():
            fields["image"] = field
            selections.append({"field": "image", "source_header": field[0],
                               "rule": "explicit_source_filename_cropped_preferred"})
            break
    return selections


def _iodp_depth_convention(entry: _Entry, fields: dict, identity: tuple) -> dict | None:
    """Apply the published CSF-A convention only to known 362 macro exports."""
    if (not re.fullmatch(r"362_macro_U\d+[A-Z]\.xlsx", PurePosixPath(entry.name).name, re.I)
            or not identity or identity[0] != "iodp" or identity[1] != "362"):
        return None
    top_header = re.sub(r" \[\d+\]$", "", _norm(fields["top"][0]))
    bottom_header = re.sub(r" \[\d+\]$", "", _norm(fields["bottom"][0]))
    if (top_header, bottom_header) != ("top depth [m]", "bottom depth [m]"):
        return None
    return {"depth_reference": "CSF-A", "rule": "iodp362_standard_sample_measurement_depth_convention",
            "source_url": IODP_METHODS_URL, "original_headers": [fields["top"][0], fields["bottom"][0]]}


def _iodp_well_from_sample(value) -> str:
    """Keep the source's explicit site/hole identifier; infer no section or depths."""
    match = re.fullmatch(r"\d+-(U\d+[A-Z])(?:-\d+[A-Z](?:-(?:\d+|CC)(?:-[AW])?)?)?",
                         str(value or "").strip())
    return match.group(1) if match else ""


def _copy_image(entry: _Entry, output_dir: Path) -> str:
    relative = f"images/{_id(entry.reference)}{PurePosixPath(entry.name).suffix.casefold()}"
    target = output_dir / relative
    if not target.exists():
        with entry.open() as source, target.open("wb") as destination:
            shutil.copyfileobj(source, destination, length=1024 * 1024)
    return relative


def _documents_and_iodp(source_id, entries, output_dir, intervals, candidates, documents, unmatched, progress, cancel):
    tables, used_images = [], set()
    photos = [e for e in entries if PurePosixPath(e.name).suffix.casefold() in IMAGE_SUFFIXES]
    images_by_name = defaultdict(list)
    for photo in photos:
        images_by_name[PurePosixPath(photo.name).name.casefold()].append(photo)
    for entry in entries:
        _check(cancel)
        suffix = PurePosixPath(entry.name).suffix.casefold()
        if suffix not in DOCUMENT_SUFFIXES:
            continue
        _progress(progress, f"Исходный документ: {PurePosixPath(entry.name).name}")
        try:
            if entry.size > MAX_DOCUMENT_BYTES:
                raise ValueError("document_size_limit")
            if suffix in {".xlsx", ".xlsm", ".csv", ".tsv"}:
                normalizations = []
                for sheet, number, raw in _table_rows(entry, normalizations):
                    _check(cancel)
                    address = f"{sheet}!{number}" if sheet else str(number)
                    base = _base_record(source_id, IODP_URL if "iodp" in source_id.casefold() else UKCS_URL, entry, address)
                    if normalizations:
                        base["source_workbook_normalizations"] = list(normalizations)
                    doc = dict(base, text=json.dumps(raw, ensure_ascii=False, default=str), original_cells=raw)
                    documents.append(doc)
                    fields = _fields(raw)
                    selections = _iodp_source_fields(entry, raw, fields) if "iodp" in source_id.casefold() else []
                    if not fields:
                        continue
                    value = lambda name: fields.get(name, ("", ""))[1]
                    top, bottom = _number(value("top")), _number(value("bottom"))
                    label, description = str(value("label") or ""), str(value("description") or "")
                    photo_name = str(value("image") or "").strip()
                    identity = _identity(fields)
                    if top is None or bottom is None or bottom <= top:
                        if label or description or photo_name:
                            unmatched.append(dict(base, reason="table_row_has_no_unambiguous_metre_depth_interval"))
                        continue
                    depth_header = fields["top"][0]
                    bottom_header = fields["bottom"][0]
                    # CSF-A/B etc must agree. Header strings remain available.
                    scales = re.findall(r"CSF-[A-Z]|CCSF-[A-Z]|m(?:bsf|cd)", depth_header + " " + bottom_header, re.I)
                    if len(set(s.casefold() for s in scales)) > 1:
                        unmatched.append(dict(base, reason="different_top_bottom_depth_scales"))
                        continue
                    well = str(value("well id") or value("well") or "".join(str(value(k) or "") for k in ("site", "hole")))
                    if not well and "iodp" in source_id.casefold():
                        well = _iodp_well_from_sample(value("sample id") or value("sample identifier"))
                    record = dict(base, well_id=well, source_group_id=f"{source_id}:{well or entry.reference}",
                                  source_depth_group_id=_id(source_id, identity or entry.reference, top, bottom),
                                  source_label=label, original_description=description, depth_top_m=top,
                                  depth_bottom_m=bottom, depth_reference=scales[0] if scales else depth_header,
                                  original_cells=raw, image="", image_reference=photo_name,
                                  pairing_status="no_verified_photo_link", pixel_alignment_verified=False)
                    if selections:
                        record["source_field_selections"] = selections
                    convention = _iodp_depth_convention(entry, fields, identity) if "iodp" in source_id.casefold() else None
                    if convention:
                        record["depth_reference"] = convention["depth_reference"]
                        record["source_depth_reference_convention"] = convention
                    tables.append((record, identity))
                    if label or description:
                        intervals.append(record)
            else:
                with entry.open() as stream:
                    data = stream.read(MAX_DOCUMENT_BYTES + 1)
                if len(data) > MAX_DOCUMENT_BYTES:
                    raise ValueError("document_size_limit")
                if suffix == ".docx":
                    from docx import Document
                    doc = Document(io.BytesIO(data))
                    text = "\n".join(p.text for p in doc.paragraphs)
                    text += "\n" + "\n".join("\t".join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
                elif suffix == ".pdf":
                    from pypdf import PdfReader
                    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
                else:
                    text = data.decode("utf-8-sig", errors="replace")
                base = _base_record(source_id, IODP_URL if "iodp" in source_id.casefold() else UKCS_URL, entry, "document")
                documents.append(dict(base, text=text))
                if not text.strip():
                    unmatched.append(dict(base, reason="no_extractable_text_no_ocr_attempted"))
        except DatasetPreparationCancelled:
            raise
        except Exception as exc:
            unmatched.append({"source_path": entry.reference, "reason": "document_not_parsed", "detail": str(exc)})
    photo_rows = []
    for record, identity in tables:
        if not record["image_reference"]:
            continue
        basename = PurePosixPath(record["image_reference"].replace("\\", "/")).name.casefold()
        matches = images_by_name.get(basename, [])
        if len(matches) != 1:
            unmatched.append(dict(record, reason="photo_reference_missing_or_ambiguous", match_count=len(matches)))
            continue
        photo = matches[0]
        record["image"] = _copy_image(photo, output_dir)
        record["image_source"] = photo.reference
        record["pairing_status"] = "explicit_source_image_and_depth_unreviewed"
        photo_rows.append((record, identity))
        used_images.add(photo.reference)
    for record, identity in tables:
        if not (record["source_label"] or record["original_description"]):
            continue
        matches = [(record, record)] if record["image"] else []
        if not matches and identity:
            for photo, photo_identity in photo_rows:
                if identity != photo_identity or _norm(record["depth_reference"]) != _norm(photo["depth_reference"]):
                    continue
                if photo["depth_top_m"] <= record["depth_top_m"] < record["depth_bottom_m"] <= photo["depth_bottom_m"]:
                    matches.append((record, photo))
        if len(matches) == 1:
            _, photo = matches[0]
            record.update(image=photo["image"], image_source=photo["image_source"],
                          pairing_status="source_identifier_and_depth_matched_unreviewed",
                          photo_source_record_id=photo["record_id"],
                          photo_source_path=photo["source_path"], photo_source_member=photo["source_member"],
                          photo_source_row=photo["source_row"],
                          photo_source_field_selections=photo.get("source_field_selections", []),
                          photo_depth_top_m=photo["depth_top_m"], photo_depth_bottom_m=photo["depth_bottom_m"],
                          needs_interval_crop=(record["depth_top_m"], record["depth_bottom_m"]) !=
                                              (photo["depth_top_m"], photo["depth_bottom_m"]))
            candidates.append({"image": record["image"], "source_label": record["source_label"],
                               "original_description": record["original_description"], "metadata": dict(record)})
        else:
            unmatched.append(dict(record, reason="interval_has_no_unique_identifier_depth_photo_match", match_count=len(matches)))
    for photo in photos:
        if photo.reference not in used_images:
            unmatched.append({"source_path": photo.reference, "reason": "photo_without_source_interval_reference"})


def _finish(source_id, entries, output_dir, intervals, candidates, documents, unmatched):
    _jsonl(output_dir / "intervals.jsonl", intervals)
    _jsonl(output_dir / "candidates.jsonl", candidates)
    _jsonl(output_dir / "source_documents.jsonl", documents)
    _jsonl(output_dir / "unmatched.jsonl", unmatched)
    fields = ["record_id", "well_id", "depth_top_m", "depth_bottom_m", "source_label", "original_description",
              "source_depth_group_id", "image", "pairing_status", "training_approved", "geologist_notes"]
    with (output_dir / "review.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for record in intervals:
            # Protect opening the review CSV in Excel from formula execution.
            writer.writerow({key: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
                             for key, value in record.items() if key in fields})
    inventory = [{"path": str(e.path), "member": e.member, "bytes": e.size} for e in entries]
    (output_dir / "source_inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    knowledge_dir = output_dir / "knowledge"
    knowledge_dir.mkdir(exist_ok=True)
    grouped_documents = defaultdict(list)
    for document in documents:
        grouped_documents[(document["source_path"], document["source_member"])].append(document)
    for (source_path, member), rows in grouped_documents.items():
        text = "Исходный материал; перевод и геологическая проверка не выполнены.\n"
        text += f"Источник: {source_path}" + (f"!{member}" if member else "") + "\n\n"
        text += "\n\n".join(f"[{row['source_row']}]\n{row['text']}" for row in rows)
        (knowledge_dir / f"source_{_id(source_path, member)}.txt").write_text(text, encoding="utf-8")
    cards = []
    for candidate in candidates:
        meta = candidate["metadata"]
        cards.append(f'<article><img loading="lazy" src="{html.escape(candidate["image"], quote=True)}">'
                     f'<p>{html.escape(str(meta.get("well_id", "")))}: {meta.get("depth_top_m")}–{meta.get("depth_bottom_m")} м</p>'
                     f'<pre>{html.escape(candidate["source_label"])}\n{html.escape(candidate["original_description"])}</pre>'
                     f'<small>{html.escape(meta["record_id"])}</small></article>')
    (output_dir / "review.html").write_text(
        '<!doctype html><html lang="ru"><meta charset="utf-8"><title>Проверка исходных интервалов</title>'
        '<style>body{font:16px system-ui;margin:24px}main{display:flex;flex-wrap:wrap;gap:16px}'
        'article{border:1px solid #aaa;padding:12px;width:320px}img{max-width:100%;max-height:360px}pre{white-space:pre-wrap}</style>'
        '<h1>Исходные интервалы: требуется проверка геолога</h1>'
        '<p>Исходные коды и тексты сохранены. Перевод и допуск к обучению не выполнены. '
        'Сопоставление по глубине не подтверждает видимость каждого признака на фото. Правки записывайте в review.csv.</p>'
        '<main>' + ''.join(cards) + '</main></html>', encoding="utf-8")
    report = {"schema": "kern-open-interval-dataset-v1", "source_id": source_id,
              "status": "review_ready" if candidates else "source_index_only",
              "output_dir": str(output_dir), "manifest": str(output_dir / "candidates.jsonl"),
              "intervals_manifest": str(output_dir / "intervals.jsonl"),
              "source_documents": str(output_dir / "source_documents.jsonl"),
              "knowledge_dir": str(knowledge_dir),
              "review_csv": str(output_dir / "review.csv"), "review_html": str(output_dir / "review.html"),
              "unmatched_manifest": str(output_dir / "unmatched.jsonl"),
              "counts": {"source_files": len(entries), "intervals": len(intervals), "paired_candidates": len(candidates),
                         "documents": len(documents), "unmatched": len(unmatched), "training_approved": 0},
              "unmatched_reasons": dict(Counter(row["reason"] for row in unmatched)),
              "readiness": {"can_review": bool(candidates or documents or intervals), "can_train": False,
                            "requires_geologist_review": True, "requires_rights_review": True},
              "notes": ["No inferred or translated geological descriptions were generated.",
                        "UKCS previews are bounded to two runs per original class per well; all runs are indexed.",
                        "Depth-array first/last values are observed rows, not inferred physical interval edges.",
                        "IODP photo links require explicit source image references; cross-table joins require matching identifiers including section half and depth reference.",
                        "Known Expedition 362 LSIMG tables use their explicit CSF-A columns and prefer the published cropped filename. Known macro section rows use the documented standard CSF-A sample convention; every selection records provenance.",
                        "PDFs use existing text only; no OCR or geological extraction from figures was attempted."]}
    (output_dir / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def prepare_interval_dataset(source_id: str, source_paths: list[Path], output_dir: Path, *, progress=None, cancel=None) -> dict:
    """Create an auditable review package from local UKCS or IODP362 material.

    ``progress`` receives a human-readable string. ``cancel`` may be a callable
    or threading.Event. Existing nonempty output is rejected to preserve edits.
    No emitted record is approved for model training.
    """
    source_id = str(source_id).strip()
    if not source_id or not any(token in source_id.casefold() for token in ("ukcs", "q204", "q205", "iodp", "362")):
        raise ValueError("Поддерживаются источники UKCS Q204/Q205 и IODP362")
    output_dir = Path(output_dir).resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"Папка уже содержит данные проверки: {output_dir}")
    _check(cancel)
    entries_unmatched = []
    entries = _inventory(source_paths, output_dir, entries_unmatched, cancel)
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "images").mkdir(exist_ok=True)
    intervals, candidates, documents = [], [], []
    if any(token in source_id.casefold() for token in ("ukcs", "q204", "q205")):
        used = _ukcs(source_id, entries, output_dir, intervals, candidates, entries_unmatched, progress, cancel)
        for entry in entries:
            if PurePosixPath(entry.name).suffix.casefold() == ".npy" and entry.reference not in used:
                entries_unmatched.append({"source_path": entry.reference, "reason": "unmatched_source_array"})
    _documents_and_iodp(source_id, entries, output_dir, intervals, candidates, documents, entries_unmatched, progress, cancel)
    for entry in entries:
        suffix = PurePosixPath(entry.name).suffix.casefold()
        if suffix not in DOCUMENT_SUFFIXES | IMAGE_SUFFIXES | {".npy"}:
            entries_unmatched.append({"source_path": entry.reference, "reason": "unsupported_source_file_indexed_only"})
    _check(cancel)
    report = _finish(source_id, entries, output_dir, intervals, candidates, documents, entries_unmatched)
    _progress(progress, f"Готово к проверке: {len(intervals)} интервалов, {len(candidates)} фото; к обучению допущено 0")
    return report
