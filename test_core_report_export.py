"""Regression test for the fixed 22-column client Excel export."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.infrastructure.core_report_export import export_core_description_report


class CoreReportExportTests(unittest.TestCase):
    def test_final_excel_preserves_template_and_only_explicit_vlm_description(self):
        wells = [{
            "name": "Р-31",
            "core_interval": (3002.0, 3004.0),
            "excel_report_inputs": {
                "field": "Уренгойское", "well_number": "Р-31", "drilling_number": "1",
                "core_top": 3002.0, "core_base": 3004.0, "stratigraphy": "К2m",
                "gis_shift": -1.0, "penetration": 2.0, "core_recovered": 1.5,
            },
            "layers": [{
                "label": "L",
                "depth_from": 3002.0,
                "depth_to": 3002.5,
                "attributes": {
                    "Название фации": "Лагуна",
                    "Название породы": "Песчаник",
                    "Цвет": "Серый",
                    "Цемент": "Карбонатный",
                },
            }],
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            export_core_description_report(path, "Тест", wells, "xlsx")
            sheet = load_workbook(path, data_only=True).active

            self.assertEqual("Краткое описание", sheet["V1"].value)
            self.assertEqual(22, sheet.max_column)
            self.assertEqual("Уренгойское", sheet["A4"].value)
            self.assertEqual("Р-31", sheet["B4"].value)
            self.assertEqual(-1.0, sheet["G4"].value)
            self.assertEqual(3002.0, sheet["M4"].value)
            self.assertEqual(3002.5, sheet["N4"].value)
            self.assertEqual(1, sheet["Q4"].value)
            self.assertEqual("Лагуна", sheet["S4"].value)
            self.assertIsNone(sheet["V4"].value)


if __name__ == "__main__":
    unittest.main()
