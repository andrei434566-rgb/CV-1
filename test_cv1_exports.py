"""CV-1 report rows retain segmentation contacts independently of VLM labels."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from app.infrastructure.core_report_export import _merged_layers, export_core_description_report


class Cv1ExportTests(unittest.TestCase):
    def layer(self, top, base, *, strict=True, record="photo-a"):
        attributes = {
            "Название фации": "Фация, предположение VLM",
            "__vlm_facies_association": "Не установлена",
            "__vlm_depositional_environment": "Не установлена",
            "__vlm_description": "Серая порода с тонкой слоистостью.",
        }
        if strict:
            attributes.update({"__geometry_only": "1", "__geometry_id": "yolo:1", "__record_id": record})
        return {"label": "Фация, предположение VLM", "depth_from": top, "depth_to": base,
                "attributes": attributes}

    def test_adjacent_strict_masks_with_identical_vlm_text_retain_both_contacts(self):
        first = self.layer(100, 101)
        second = self.layer(101, 102)
        well = {"layers": [first, second]}
        original = deepcopy(well)
        rows = _merged_layers(well)
        self.assertEqual([(r["depth_from"], r["depth_to"]) for r in rows], [(100, 101), (101, 102)])
        self.assertEqual(well, original)

    def test_overlapping_strict_masks_are_not_silently_unioned(self):
        rows = _merged_layers({"layers": [self.layer(100, 101.2), self.layer(101, 102)]})
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["depth_to"], 101.2)
        self.assertEqual(rows[1]["depth_from"], 101)

    def test_per_photo_geometry_id_collision_does_not_merge_rows(self):
        rows = _merged_layers({"layers": [self.layer(100, 101, record="photo-a"),
                                          self.layer(101, 102, record="photo-b")]})
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["attributes"]["__record_id"] for r in rows], ["photo-a", "photo-b"])

    def test_mixed_legacy_and_strict_rows_do_not_merge_in_either_order(self):
        for strict_first in (True, False):
            with self.subTest(strict_first=strict_first):
                rows = _merged_layers({"layers": [self.layer(100, 101, strict=strict_first),
                                                  self.layer(101, 102, strict=not strict_first)]})
                self.assertEqual(len(rows), 2)

    def test_legacy_adjacent_facies_merge_is_preserved(self):
        rows = _merged_layers({"layers": [self.layer(100, 101, strict=False),
                                          self.layer(101, 102, strict=False)]})
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["depth_from"], rows[0]["depth_to"]), (100, 102))

    def test_actual_final_excel_retains_two_rows_and_vlm_descriptions(self):
        well = {"name": "Проверка CV-1", "layers": [self.layer(100, 101), self.layer(101, 102)]}
        temp_root = Path(__file__).resolve().parent / "open_datasets/test_tmp"
        temp_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="cv1_export_", dir=temp_root) as folder:
            target = Path(folder) / "report.xlsx"
            export_core_description_report(target, "CV-1", [well], "xlsx")
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.max_column, 22)
                self.assertEqual([(sheet[f"M{r}"].value, sheet[f"N{r}"].value) for r in (4, 5)],
                                 [(100, 101), (101, 102)])
                self.assertEqual([sheet[f"Q{r}"].value for r in (4, 5)], [1, 2])
                self.assertEqual([sheet[f"V{r}"].value for r in (4, 5)],
                                 ["Серая порода с тонкой слоистостью."] * 2)
                self.assertEqual(sheet["R4"].value, '=IF(COUNT(M4:N4)=2,N4-M4,"")')
                self.assertEqual(sheet["R5"].value, '=IF(COUNT(M5:N5)=2,N5-M5,"")')
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
