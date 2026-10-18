"""One-window Windows trainer for preparing Excel/photo pairs and fine-tuning Qwen3-VL."""

from __future__ import annotations

import csv
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import traceback
from pathlib import Path
from tkinter import BooleanVar, END, StringVar, filedialog, messagebox, ttk
import tkinter as tk

from openpyxl import load_workbook


APP_NAME = "Локальный ИИ-седиментолог — обучение"
APP_VERSION = "0.1.0"
MODEL_ID = "Qwen/Qwen3-VL-2B-Instruct"
PYTHON_VERSION = "3.11"
TRAINING_SCRIPT = "train_qwen3vl.py"
PHOTO_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
DEPTH_RE = re.compile(r"(?<!\d)(\d{3,5}(?:[.,]\d{1,3})?)\s*[-–—]\s*(\d{3,5}(?:[.,]\d{1,3})?)(?!\d)")
UV_RE = re.compile(r"(?:_uv|_уф)(?:\.|$)", re.IGNORECASE)
REQUIREMENTS = (
    "torch==2.6.0+cu126",
    "torchvision==0.21.0+cu126",
    "transformers==5.18.0",
    "peft==0.21.2",
    "accelerate==1.15.0",
    "Pillow>=10.0.0",
    "openpyxl>=3.1.0",
)


def install_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parent


def app_data_root() -> Path:
    documents = Path.home() / "Documents"
    if not documents.is_dir():
        documents = Path.home()
    return documents / "CoreSedimentologist"


def parse_depth(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        match = DEPTH_RE.search(text)
        if match:
            return float(match.group(1).replace(",", "."))
    return None


def hash_short(value: str) -> str:
    import hashlib
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:8]


class CoreTrainerApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title(APP_NAME)
        self.root.geometry("1040x760")
        self.root.minsize(880, 660)
        self.root.configure(bg="#eef2ef")
        self.events: queue.Queue[tuple[str, str]] = queue.Queue()
        self.workspace = StringVar(value=str(app_data_root()))
        self.workbook = StringVar()
        self.photos_dir = StringVar()
        self.sheet = StringVar()
        self.header_row = StringVar(value="1")
        self.top_column = StringVar()
        self.base_column = StringVar()
        self.description_column = StringVar()
        self.environment_column = StringVar(value="— нет —")
        self.epochs = StringVar(value="2")
        self.learning_rate = StringVar(value="0.00001")
        self.approved = BooleanVar(value=False)
        self.status = StringVar(value="Выберите Excel с интервалами и папку с фото.")
        self.columns: list[tuple[str, int]] = []
        self.prepared_manifest: Path | None = None
        self.prepared_count = 0
        self.language_block = False
        self.busy = False
        self._make_ui()
        self._poll_events()

    def _make_ui(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("vista")
        except tk.TclError:
            pass
        style.configure("TFrame", background="#eef2ef")
        style.configure("Card.TLabelframe", background="#ffffff", padding=12)
        style.configure("Card.TLabelframe.Label", background="#ffffff", font=("Segoe UI", 10, "bold"))
        style.configure("TLabel", background="#eef2ef", font=("Segoe UI", 9))
        style.configure("Title.TLabel", font=("Segoe UI", 18, "bold"), foreground="#163b38")
        style.configure("Sub.TLabel", foreground="#425955")
        style.configure("Action.TButton", font=("Segoe UI", 10, "bold"), padding=8)

        outer = ttk.Frame(self.root, padding=18)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text="Обучение Qwen на фото и Excel", style="Title.TLabel").pack(anchor="w")
        ttk.Label(outer, text="Подготовка пар, проверка компьютера и запуск обучения через одно окно.", style="Sub.TLabel").pack(anchor="w", pady=(2, 12))

        data_card = ttk.LabelFrame(outer, text="1. Укажите ваши данные", style="Card.TLabelframe")
        data_card.pack(fill="x", pady=6)
        self._path_row(data_card, "Excel с описаниями:", self.workbook, self._choose_workbook, 0)
        self._path_row(data_card, "Папка с фотографиями:", self.photos_dir, self._choose_photos, 1)
        self._path_row(data_card, "Папка результатов:", self.workspace, self._choose_workspace, 2)

        mapping_card = ttk.LabelFrame(outer, text="2. Сопоставьте столбцы Excel", style="Card.TLabelframe")
        mapping_card.pack(fill="x", pady=6)
        row = ttk.Frame(mapping_card)
        row.pack(fill="x", pady=3)
        ttk.Label(row, text="Лист").pack(side="left")
        self.sheet_box = ttk.Combobox(row, textvariable=self.sheet, state="readonly", width=24)
        self.sheet_box.pack(side="left", padx=(6, 14))
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _e: self._load_columns())
        ttk.Label(row, text="Строка заголовков").pack(side="left")
        ttk.Entry(row, textvariable=self.header_row, width=6).pack(side="left", padx=6)
        ttk.Button(row, text="Прочитать заголовки", command=self._load_columns).pack(side="left", padx=4)
        fields = ttk.Frame(mapping_card)
        fields.pack(fill="x", pady=4)
        self.top_box = self._column_box(fields, "Верх интервала", self.top_column, 0)
        self.base_box = self._column_box(fields, "Низ интервала", self.base_column, 1)
        self.description_box = self._column_box(fields, "Описание породы", self.description_column, 2)
        self.environment_box = self._column_box(fields, "Обстановка (если есть)", self.environment_column, 3)
        ttk.Label(mapping_card, text="Фото сопоставляются с интервалами по глубинам в имени файла, например 3841.00-3841.98.", style="Sub.TLabel").pack(anchor="w", pady=(4, 0))

        review_card = ttk.LabelFrame(outer, text="3. Подтвердите привязку", style="Card.TLabelframe")
        review_card.pack(fill="x", pady=6)
        self.prepare_button = ttk.Button(review_card, text="Подготовить пары и таблицу для проверки", style="Action.TButton", command=self._prepare_data)
        self.prepare_button.pack(side="left", padx=(0, 12))
        self.open_preview_button = ttk.Button(review_card, text="Открыть таблицу пар", command=self._open_preview, state="disabled")
        self.open_preview_button.pack(side="left", padx=(0, 12))
        self.approval_check = ttk.Checkbutton(review_card, text="Я проверил соответствие фото, интервалов и описаний; разрешаю использовать эти пары для обучения", variable=self.approved, command=self._approval_changed)
        self.approval_check.pack(side="left", fill="x", expand=True)

        train_card = ttk.LabelFrame(outer, text="4. Запустите обучение", style="Card.TLabelframe")
        train_card.pack(fill="x", pady=6)
        row = ttk.Frame(train_card)
        row.pack(fill="x")
        ttk.Label(row, text="Эпохи").pack(side="left")
        ttk.Spinbox(row, from_=1, to=10, textvariable=self.epochs, width=5).pack(side="left", padx=(6, 18))
        ttk.Label(row, text="Learning rate").pack(side="left")
        ttk.Entry(row, textvariable=self.learning_rate, width=12).pack(side="left", padx=6)
        self.train_button = ttk.Button(row, text="Проверить ПК и обучить Qwen", style="Action.TButton", command=self._start_training, state="disabled")
        self.train_button.pack(side="right")
        ttk.Label(train_card, text="Первый запуск загрузит Python, CUDA-библиотеки и Qwen 2B (несколько ГБ). Нужны интернет, Windows x64, драйвер NVIDIA, видеокарта с 8 ГБ VRAM и примерно 25 ГБ свободного места. Тексты берутся из Excel без перевода: для русскоязычной модели выберите русские описания.", style="Sub.TLabel", wraplength=950).pack(anchor="w", pady=(9, 0))

        status_bar = ttk.Frame(outer)
        status_bar.pack(fill="x", pady=(8, 4))
        ttk.Label(status_bar, textvariable=self.status, style="Sub.TLabel").pack(side="left", fill="x", expand=True)
        self.progress = ttk.Progressbar(status_bar, mode="indeterminate", length=150)
        self.progress.pack(side="right")
        log_card = ttk.LabelFrame(outer, text="Ход работы", style="Card.TLabelframe")
        log_card.pack(fill="both", expand=True, pady=4)
        self.log = tk.Text(log_card, height=10, wrap="word", font=("Consolas", 9), bg="#101a1c", fg="#edf4ee", insertbackground="#ffffff", relief="flat")
        scroll = ttk.Scrollbar(log_card, command=self.log.yview)
        self.log.configure(yscrollcommand=scroll.set)
        self.log.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        self._log("Готово. Выберите исходный Excel и папку с фото.")

    def _path_row(self, parent: ttk.Widget, label: str, variable: StringVar, command, row: int) -> None:
        frame = ttk.Frame(parent)
        frame.pack(fill="x", pady=3)
        ttk.Label(frame, text=label, width=24).pack(side="left")
        ttk.Entry(frame, textvariable=variable).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(frame, text="Обзор…", command=command).pack(side="right")

    def _column_box(self, parent: ttk.Widget, title: str, variable: StringVar, column: int) -> ttk.Combobox:
        frame = ttk.Frame(parent)
        frame.grid(row=0, column=column, sticky="ew", padx=4)
        parent.columnconfigure(column, weight=1)
        ttk.Label(frame, text=title).pack(anchor="w")
        box = ttk.Combobox(frame, textvariable=variable, state="readonly")
        box.pack(fill="x", pady=(3, 0))
        return box

    def _choose_workbook(self) -> None:
        path = filedialog.askopenfilename(title="Выберите Excel с интервальными описаниями", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not path:
            return
        self.workbook.set(path)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            self.sheet_box["values"] = wb.sheetnames
            self.sheet.set(wb.sheetnames[0])
            wb.close()
            self._load_columns()
        except Exception as exc:
            messagebox.showerror("Не удалось прочитать Excel", str(exc))

    def _choose_photos(self) -> None:
        path = filedialog.askdirectory(title="Выберите папку с фотографиями керна")
        if path:
            self.photos_dir.set(path)

    def _choose_workspace(self) -> None:
        path = filedialog.askdirectory(title="Выберите папку, где сохранять обучение и модель")
        if path:
            self.workspace.set(path)

    def _load_columns(self) -> None:
        if not self.workbook.get() or not self.sheet.get():
            return
        try:
            header = int(self.header_row.get())
            if header < 1:
                raise ValueError("Строка заголовка должна быть не меньше 1")
            wb = load_workbook(self.workbook.get(), read_only=True, data_only=True)
            ws = wb[self.sheet.get()]
            first = next(ws.iter_rows(min_row=header, max_row=header, values_only=True), ())
            sample = next(ws.iter_rows(min_row=header + 1, max_row=header + 1, values_only=True), ())
            values: list[tuple[str, int]] = []
            for idx, item in enumerate(first):
                if item is None and idx >= len(sample):
                    continue
                letter = self._excel_column(idx + 1)
                label = str(item).strip() if item is not None else ""
                example = str(sample[idx]).strip() if idx < len(sample) and sample[idx] is not None else ""
                display = f"{letter} — {label or '(без заголовка)'}" + (f" · {example[:45]}" if example else "")
                values.append((display, idx))
            wb.close()
            self.columns = values
            choices = [v[0] for v in values]
            for box in (self.top_box, self.base_box, self.description_box):
                box["values"] = choices
            self.environment_box["values"] = ["— нет —", *choices]
            if choices:
                self._auto_select(self.top_box, ("верх", "от", "top", "начало"))
                self._auto_select(self.base_box, ("низ", "до", "base", "конец"))
                self._auto_select(self.description_box, ("опис", "литолог", "характеристика", "description"))
            self._log(f"Прочитано столбцов: {len(values)}. Примеры значений показаны в списках.")
        except Exception as exc:
            messagebox.showerror("Ошибка Excel", str(exc))

    @staticmethod
    def _excel_column(number: int) -> str:
        out = ""
        while number:
            number, rem = divmod(number - 1, 26)
            out = chr(65 + rem) + out
        return out

    def _auto_select(self, box: ttk.Combobox, words: tuple[str, ...]) -> None:
        for label, _idx in self.columns:
            if any(word in label.lower() for word in words):
                box.set(label)
                break

    def _column_index(self, box: ttk.Combobox) -> int | None:
        value = box.get()
        return next((idx for label, idx in self.columns if label == value), None)

    def _prepare_data(self) -> None:
        if (not self.workbook.get() or not self.photos_dir.get() or not self.sheet.get()
                or self._column_index(self.top_box) is None
                or self._column_index(self.base_box) is None
                or self._column_index(self.description_box) is None):
            messagebox.showwarning("Не заполнены поля", "Выберите Excel, лист, фото и столбцы глубин и описания.")
            return
        self.approved.set(False)
        self._approval_changed()
        config = {
            "workbook": self.workbook.get(), "photos_dir": self.photos_dir.get(), "workspace": self.workspace.get(),
            "sheet": self.sheet.get(), "header": int(self.header_row.get()), "top_col": self._column_index(self.top_box),
            "base_col": self._column_index(self.base_box), "desc_col": self._column_index(self.description_box),
            "env_col": self._column_index(self.environment_box) if self.environment_column.get() != "— нет —" else None,
            "approved": bool(self.approved.get()),
        }
        self._set_busy(True, "Сопоставляю фото с интервалами Excel…")
        thread = threading.Thread(target=self._prepare_worker, args=(config,), daemon=True)
        thread.start()

    def _prepare_worker(self, config: dict) -> None:
        try:
            workdir = Path(config["workspace"]).expanduser().resolve()
            data_dir = workdir / "data"
            data_dir.mkdir(parents=True, exist_ok=True)
            wb = load_workbook(config["workbook"], read_only=True, data_only=True)
            ws = wb[config["sheet"]]
            header = config["header"]
            top_col = config["top_col"]
            base_col = config["base_col"]
            desc_col = config["desc_col"]
            env_col = config["env_col"]
            assert top_col is not None and base_col is not None and desc_col is not None
            layers: list[dict] = []
            for excel_row, values in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), start=header + 1):
                if max(top_col, base_col, desc_col) >= len(values):
                    continue
                top = parse_depth(values[top_col])
                base = parse_depth(values[base_col])
                description = " ".join(str(values[desc_col] or "").split())
                environment = ""
                if env_col is not None and env_col < len(values):
                    environment = " ".join(str(values[env_col] or "").split())
                if top is None or base is None or base <= top or len(description) < 8:
                    continue
                layers.append({"row": excel_row, "top": top, "base": base, "description": description, "environment": environment})
            wb.close()
            if not layers:
                raise ValueError("В выбранных столбцах не найдены строки с глубиной и описанием. Проверьте строку заголовка и столбцы.")

            image_paths = sorted(p for p in Path(config["photos_dir"]).rglob("*") if p.is_file() and p.suffix.lower() in PHOTO_EXTENSIONS and not UV_RE.search(p.name))
            if not image_paths:
                raise ValueError("В выбранной папке не найдены поддерживаемые фото.")
            rows: list[dict] = []
            pair_preview: list[list[str]] = []
            no_depth = 0
            no_match = 0
            for image_path in image_paths:
                match = DEPTH_RE.search(image_path.stem)
                if not match:
                    no_depth += 1
                    continue
                photo_top, photo_base = (float(v.replace(",", ".")) for v in match.groups())
                if photo_base <= photo_top:
                    continue
                segments = []
                for layer in layers:
                    left, right = max(photo_top, layer["top"]), min(photo_base, layer["base"])
                    if right <= left:
                        continue
                    segments.append({
                        "top_m": round(left, 3), "base_m": round(right, 3), "description": layer["description"],
                        "source_row": layer["row"], "environment": layer["environment"] or None,
                    })
                spans = sorted((seg["top_m"], seg["base_m"]) for seg in segments)
                covered = 0.0
                span_start, span_end = spans[0]
                for left, right in spans[1:]:
                    if left <= span_end:
                        span_end = max(span_end, right)
                    else:
                        covered += span_end - span_start
                        span_start, span_end = left, right
                covered += span_end - span_start
                coverage = min(1.0, covered / (photo_base - photo_top))
                if not segments or coverage < 0.80:
                    no_match += 1
                    continue
                rid = f"{image_path.stem}_{hash_short(str(image_path))}"
                target: dict = {"description_segments": [{k: s[k] for k in ("top_m", "base_m", "description")} for s in segments]}
                env_targets = list(dict.fromkeys(s["environment"] for s in segments if s["environment"]))
                if env_targets:
                    target["environment"] = env_targets
                prompt = (
                    f"<image>\nОпиши керн на фотографии в пределах {photo_top:.3f}–{photo_base:.3f} м. "
                    "Укажи видимые цвет, зернистость, текстуру, слоистость и контакты; для описанных интервалов верни JSON "
                    "с ключом description_segments и объектами top_m, base_m, description. "
                    "Не добавляй свойства, которых нет на фото или в проверенном описании. "
                    "Обстановку указывай только если для интервала есть отдельная проверенная метка."
                )
                rows.append({
                    "image": str(image_path.resolve()),
                    "conversations": [
                        {"from": "human", "value": prompt},
                        {"from": "gpt", "value": json.dumps(target, ensure_ascii=False)},
                    ],
                    "metadata": {
                        "record_id": rid, "well_id": "", "photo_depth_interval_m": [photo_top, photo_base],
                        "source_excel": str(Path(config["workbook"]).resolve()), "sheet": config["sheet"],
                        "description_row_numbers": [s["source_row"] for s in segments],
                        "coverage_fraction": round(coverage, 4), "environment_labels_present": bool(env_targets),
                        "training_approved": config["approved"], "geology_approved": False,
                        "approval_basis": "user confirmed image/interval/description pairing in GUI" if config["approved"] else "pending user confirmation in GUI",
                    },
                })
                pair_preview.append([rid, str(image_path), f"{photo_top:.3f}", f"{photo_base:.3f}", str(len(segments)), f"{coverage:.3f}", " | ".join(s["description"] for s in segments)])
            if not rows:
                raise ValueError(f"Не найдено пар. Фото без глубин в имени: {no_depth}; без подходящего описания: {no_match}. Ожидается интервал в имени вроде 3841.00-3841.98.")

            manifest = data_dir / "training_manifest.jsonl"
            manifest.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
            preview = data_dir / "pairs_to_review.csv"
            with preview.open("w", encoding="utf-8-sig", newline="") as csv_file:
                writer = csv.writer(csv_file, delimiter=";")
                writer.writerow(["record_id", "photo", "top_m", "base_m", "segments", "coverage", "source_description"])
                writer.writerows(pair_preview)
            summary = {
                "manifest": str(manifest), "preview_csv": str(preview), "matched_photos": len(rows),
                "source_description_rows": len({n for row in rows for n in row["metadata"]["description_row_numbers"]}),
                "photos_missing_depth": no_depth, "photos_without_match": no_match,
                "target_text_mostly_latin": self._target_text_mostly_latin(rows),
                "training_approved": config["approved"], "note": "Matched by filename depth interval against selected Excel depth columns; review pairs_to_review.csv before training.",
            }
            (data_dir / "pairing_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
            self.events.put(("prepared", json.dumps(summary, ensure_ascii=False)))
        except Exception:
            self.events.put(("error", traceback.format_exc()))

    @staticmethod
    def _target_text_mostly_latin(rows: list[dict]) -> bool:
        text_parts = []
        for row in rows:
            target = json.loads(row["conversations"][1]["value"])
            text_parts.extend(segment.get("description", "") for segment in target.get("description_segments", []))
            text_parts.extend(target.get("environment", []))
        text = " ".join(text_parts)
        cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
        latin = len(re.findall(r"[A-Za-z]", text))
        return latin >= 40 and latin > cyrillic

    def _start_training(self) -> None:
        if not self.prepared_manifest or self.prepared_count < 1:
            messagebox.showwarning("Сначала подготовьте пары", "Нажмите «Подготовить пары» и проверьте таблицу сопоставлений.")
            return
        if self.language_block:
            messagebox.showwarning("Нужны русские учебные описания", "В выбранном столбце большинство целевого текста латиницей. Эта сборка готовит обучение с ответами на русском и не переводит данные автоматически. Выберите русский столбец описаний или подготовьте проверенный перевод, затем заново сопоставьте пары.")
            return
        if not self.approved.get():
            messagebox.showwarning("Нужно подтверждение", "Сначала проверьте таблицу пар и отметьте, что фото, глубины и описания соответствуют друг другу.")
            return
        try:
            epochs = int(self.epochs.get())
            learning_rate = float(self.learning_rate.get())
            if epochs < 1 or not (0 < learning_rate < 1):
                raise ValueError
        except ValueError:
            messagebox.showerror("Параметры обучения", "Эпохи должны быть целым числом от 1, learning rate — между 0 и 1.")
            return
        try:
            rows = [json.loads(line) for line in self.prepared_manifest.read_text(encoding="utf-8").splitlines() if line]
            for row in rows:
                row.setdefault("metadata", {})["training_approved"] = True
                row["metadata"]["approval_basis"] = "user confirmed image/interval/description pairing in GUI"
            self.prepared_manifest.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        except Exception as exc:
            messagebox.showerror("Ошибка манифеста", str(exc))
            return
        config = {"workspace": self.workspace.get(), "manifest": str(self.prepared_manifest), "epochs": epochs, "learning_rate": learning_rate}
        self._set_busy(True, "Проверяю GPU и готовлю окружение…")
        thread = threading.Thread(target=self._train_worker, args=(config,), daemon=True)
        thread.start()

    def _find_uv(self) -> Path:
        candidates = [install_root() / "uv.exe", Path(sys.executable).parent / "uv.exe"]
        for candidate in candidates:
            if candidate.is_file():
                return candidate
        raise FileNotFoundError("В приложении не найден менеджер Python uv.exe.")

    def _run_process(self, command: list[str], cwd: Path, env: dict[str, str], label: str) -> str:
        self.events.put(("status", label))
        self.events.put(("log", "$ " + subprocess.list2cmdline(command)))
        process = subprocess.Popen(command, cwd=cwd, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", bufsize=1, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        assert process.stdout is not None
        lines = []
        for line in process.stdout:
            self.events.put(("log", line.rstrip()))
            lines.append(line.rstrip())
        code = process.wait()
        if code != 0:
            raise RuntimeError(f"Этап «{label}» завершился с кодом {code}. См. журнал выше.")
        return "\n".join(lines)

    def _train_worker(self, config: dict) -> None:
        try:
            epochs = config["epochs"]
            learning_rate = config["learning_rate"]
            workspace = Path(config["workspace"]).expanduser().resolve()
            workspace.mkdir(parents=True, exist_ok=True)
            free_gb = shutil.disk_usage(workspace).free / (1024 ** 3)
            if free_gb < 25:
                raise RuntimeError(f"На выбранном диске свободно {free_gb:.1f} ГБ. Нужно не менее 25 ГБ для Python, CUDA-библиотек, модели и результатов.")
            gpu = shutil.which("nvidia-smi")
            if not gpu:
                raise RuntimeError("Не найден nvidia-smi. Установите драйвер NVIDIA и перезапустите приложение.")
            gpu_output = self._run_process([gpu, "--query-gpu=name,memory.total,driver_version", "--format=csv,noheader"], workspace, os.environ.copy(), "Проверяю видеокарту NVIDIA…")
            self.events.put(("log", "Обнаружен GPU: " + gpu_output.replace("\n", " | ")))
            memory_match = re.search(r"(?:^|,\s*)(\d+)\s*MiB(?:,|$)", gpu_output, re.IGNORECASE)
            if memory_match and int(memory_match.group(1)) < 7680:
                raise RuntimeError(f"У видеокарты только {int(memory_match.group(1))} MiB памяти. Для этой сборки нужно не менее 8 ГБ VRAM; найдено меньше.")

            uv = self._find_uv()
            runtime = workspace / ".runtime"
            env_dir = runtime / ".venv"
            runtime.mkdir(parents=True, exist_ok=True)
            app_env = os.environ.copy()
            app_env["HF_HOME"] = str(workspace / "model_cache")
            app_env["UV_CACHE_DIR"] = str(workspace / ".uv_cache")
            app_env["UV_PYTHON_INSTALL_DIR"] = str(workspace / ".runtime_python")
            app_env["HF_HUB_DISABLE_TELEMETRY"] = "1"
            app_env["TOKENIZERS_PARALLELISM"] = "false"
            py = env_dir / "Scripts" / "python.exe"
            if not py.exists():
                self._run_process([str(uv), "python", "install", PYTHON_VERSION], workspace, app_env, "Скачиваю Python 3.11…")
                self._run_process([str(uv), "venv", "--python", PYTHON_VERSION, str(env_dir)], workspace, app_env, "Создаю отдельное окружение…")
            import_site_packages = [str(uv), "pip", "install", "--python", str(py), "--extra-index-url", "https://download.pytorch.org/whl/cu126", *REQUIREMENTS]
            probe = [str(py), "-c", "import torch,transformers,peft,accelerate,openpyxl; print('Python',__import__('sys').version.split()[0]); print('Torch',torch.__version__,'CUDA wheel',torch.version.cuda,'GPU visible',torch.cuda.is_available()); print('CUDA_OK=True' if torch.cuda.is_available() else 'CUDA_OK=False'); print('Transformers',transformers.__version__)"]
            try:
                probe_output = self._run_process(probe, workspace, app_env, "Проверяю окружение…")
            except RuntimeError:
                self._run_process(import_site_packages, workspace, app_env, "Скачиваю CUDA-библиотеки и зависимости обучения…")
                probe_output = self._run_process(probe, workspace, app_env, "Проверяю установленное окружение…")
            if "CUDA_OK=True" not in probe_output:
                raise RuntimeError("PyTorch не видит CUDA-видеокарту. Проверьте драйвер NVIDIA и совместимость GPU.")
            if not (install_root() / TRAINING_SCRIPT).is_file():
                raise FileNotFoundError(f"В пакете отсутствует {TRAINING_SCRIPT}.")
            script_target = runtime / TRAINING_SCRIPT
            if install_root() != runtime:
                shutil.copy2(install_root() / TRAINING_SCRIPT, script_target)
            output_model = workspace / "models" / "core_vlm_custom_lora"
            args = [
                str(py), str(script_target), "--model", MODEL_ID,
                "--manifest", config["manifest"],
                "--output-dir", str(output_model),
                "--epochs", str(epochs), "--learning-rate", str(learning_rate),
                "--max-image-side", "448",
            ]
            self._run_process(args, workspace, app_env, "Обучаю LoRA Qwen3-VL…")
            self.events.put(("done", str(output_model)))
        except Exception:
            self.events.put(("error", traceback.format_exc()))

    def _open_preview(self) -> None:
        if self.prepared_manifest:
            preview = self.prepared_manifest.parent / "pairs_to_review.csv"
            if preview.is_file():
                os.startfile(preview)  # type: ignore[attr-defined]

    def _set_busy(self, busy: bool, status: str) -> None:
        self.busy = busy
        self.status.set(status)
        if busy:
            self.progress.start(12)
            self.prepare_button.configure(state="disabled")
            self.approval_check.configure(state="disabled")
            self.train_button.configure(state="disabled")
        else:
            self.progress.stop()
            self.prepare_button.configure(state="normal")
            self.approval_check.configure(state="normal")
            self.train_button.configure(state="normal" if self.prepared_count and self.approved.get() and not self.language_block else "disabled")

    def _approval_changed(self) -> None:
        enabled = bool(self.prepared_count and self.approved.get() and not self.language_block and not self.busy)
        self.train_button.configure(state="normal" if enabled else "disabled")

    def _log(self, text: str) -> None:
        self.log.insert(END, text + "\n")
        self.log.see(END)

    def _poll_events(self) -> None:
        try:
            while True:
                event, text = self.events.get_nowait()
                if event == "log":
                    self._log(text)
                elif event == "status":
                    self.status.set(text)
                elif event == "prepared":
                    summary = json.loads(text)
                    self.prepared_manifest = Path(summary["manifest"])
                    self.prepared_count = summary["matched_photos"]
                    self.language_block = bool(summary.get("target_text_mostly_latin"))
                    self.open_preview_button.configure(state="normal")
                    if self.approved.get() and not self.language_block:
                        self.train_button.configure(state="normal")
                    self.status.set(f"Сопоставлено фото: {self.prepared_count}; интервалов описания: {summary['source_description_rows']}.")
                    self._log(json.dumps(summary, ensure_ascii=False, indent=2))
                    self._set_busy(False, self.status.get())
                    if self.language_block:
                        messagebox.showwarning("Нужны русские учебные описания", "В выбранном столбце большинство целевого текста латиницей. Эта сборка готовит обучение с ответами на русском и не переводит данные автоматически. Выберите русский столбец описаний или подготовьте проверенный перевод, затем заново сопоставьте пары.")
                elif event == "done":
                    self._set_busy(False, "Обучение завершено.")
                    self._log(f"Адаптер сохранён: {text}")
                    messagebox.showinfo("Готово", f"Обучение завершено. Адаптер и training_metrics.json:\n{text}")
                elif event == "error":
                    self._set_busy(False, "Ошибка. См. журнал.")
                    self._log(text)
                    messagebox.showerror("Не удалось продолжить", text[-3500:])
        except queue.Empty:
            pass
        self.root.after(150, self._poll_events)


def main() -> None:
    root = tk.Tk()
    CoreTrainerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
