from __future__ import annotations

import json
import hashlib
import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

CORE_ROOT = Path(r"E:\Керн фото")
PROJECT_ROOT = Path(__file__).resolve().parent
ARCHIVE = CORE_ROOT / "!извлеченные" / "518431.zip"
WORKBOOK = CORE_ROOT / "4" / "DVD_Отчёт" / "Приложения" / "Прил.8 Описание керна_41.xlsx"
PHOTO_INTERVAL_RE = re.compile(r"(\d{4}[.,]\d{1,2})\s*[-–]\s*(\d{4}[.,]\d{1,2})")
PHOTO_NAME_RE = re.compile(r"^(?:core_day_)?[рp]-41_\d{4}[.,]\d{1,2}-\d{4}[.,]\d{1,2}\.jpe?g$", re.I)
UV_RE = re.compile(r"(?:_uv|_уф)\.jpe?g$", re.I)
MIN_COVERAGE = 0.80


def depth_number(value):
    try:
        return float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None


def union_length(intervals):
    if not intervals:
        return 0.0
    spans = sorted(intervals)
    start, end = spans[0]
    total = 0.0
    for left, right in spans[1:]:
        if left <= end:
            end = max(end, right)
        else:
            total += end - start
            start, end = left, right
    return total + end - start


def read_layers():
    ws = load_workbook(WORKBOOK, read_only=True, data_only=True).worksheets[0]
    layers = []
    # In this report the tied layer top/base are columns P/Q and the lithology text is R.
    for row_no, row in enumerate(ws.iter_rows(min_row=7, values_only=True), start=7):
        top, base = depth_number(row[15]), depth_number(row[16])
        text = str(row[17] or "").strip()
        formation = str(row[3] or "").strip()
        if top is None or base is None or base <= top or not text or text in {"-", "—"}:
            continue
        layers.append({
            "source_row": row_no,
            "top_m": top,
            "base_m": base,
            "formation": formation,
            "description": text,
        })
    return layers


def build_pairs():
    if not ARCHIVE.is_file() or not WORKBOOK.is_file():
        raise FileNotFoundError("Expected source archive or interval description workbook is missing")
    layers = read_layers()
    candidates = []
    omitted = []
    daylight_count = 0
    with zipfile.ZipFile(ARCHIVE) as zf:
        for member in zf.namelist():
            base_name = Path(member).name
            normalized = member.casefold().replace("\\", "/")
            if "/прил.3 фото керна/" not in normalized:
                continue
            if not PHOTO_NAME_RE.fullmatch(base_name) or UV_RE.search(base_name):
                continue
            daylight_count += 1
            match = PHOTO_INTERVAL_RE.search(base_name)
            if not match:
                omitted.append({"image_member": member, "reason": "depth_interval_missing_from_filename"})
                continue
            image_top, image_base = (float(s.replace(",", ".")) for s in match.groups())
            image_span = image_base - image_top
            if image_span <= 0:
                omitted.append({"image_member": member, "reason": "invalid_depth_interval"})
                continue

            segments = []
            covered_spans = []
            overlap_by_layer = []
            for layer in layers:
                left = max(image_top, layer["top_m"])
                right = min(image_base, layer["base_m"])
                if right <= left:
                    continue
                covered_spans.append((left, right))
                segment = {
                    "top_m": round(left, 3),
                    "base_m": round(right, 3),
                    "formation": layer["formation"],
                    "source_description": layer["description"],
                    "source_row": layer["source_row"],
                }
                segments.append(segment)
                overlap_by_layer.append(right - left)
            coverage = min(1.0, union_length(covered_spans) / image_span)
            if coverage < MIN_COVERAGE or not segments:
                omitted.append({
                    "image_member": member,
                    "image_top_m": image_top,
                    "image_base_m": image_base,
                    "description_coverage": round(coverage, 4),
                    "reason": "insufficient_overlap_with_described_layer",
                })
                continue

            best_layer_fraction = max(overlap_by_layer) / image_span
            match_type = "single_layer_dominant" if best_layer_fraction >= 0.65 else "multiple_described_layers"
            formations = list(dict.fromkeys(s["formation"] for s in segments if s["formation"]))
            member_tag = hashlib.sha1(member.encode("utf-8")).hexdigest()[:8]
            record_id = f"urengoy_41_{image_top:.2f}_{image_base:.2f}_{member_tag}".replace(".", "p")
            candidates.append({
                "id": record_id,
                "image": {
                    "archive": str(ARCHIVE),
                    "member": member,
                },
                "well": {"field": "Уренгойское", "number": "41"},
                "image_depth_interval_m": {"top": image_top, "base": image_base},
                "formation_candidates": formations,
                "source_description_segments": segments,
                "description_coverage_fraction": round(coverage, 4),
                "pair_type": match_type,
                "environment_label": None,
                "split": None,
                "review_status": "candidate_requires_geologist_review",
                "training_approved": False,
                "provenance": {
                    "description_workbook": str(WORKBOOK),
                    "description_sheet": "Пр.8",
                    "matching_rule": "image interval from filename intersected with tied layer top/base in workbook columns P/Q; descriptions from column R",
                },
            })
    return layers, candidates, omitted, daylight_count


def main():
    layers, candidates, omitted, daylight_count = build_pairs()
    image_dir = PROJECT_ROOT / "data" / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    qwen_output = PROJECT_ROOT / "data" / "qwen3_vl_candidates.jsonl"
    extracted_bytes = 0
    with zipfile.ZipFile(ARCHIVE) as zf:
        for record in candidates:
            image_path = image_dir / f"{record['id']}.jpg"
            payload = zf.read(record["image"]["member"])
            image_path.write_bytes(payload)
            extracted_bytes += len(payload)
            record["image"]["local_copy"] = str(image_path)

    output = PROJECT_ROOT / "data" / "candidate_pairs.jsonl"
    with output.open("w", encoding="utf-8", newline="\n") as f:
        for record in candidates:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")

    # Qwen3-VL's official fine-tuning data shape: image + human/gpt conversations.
    # Keep the exact source wording and mark every row as a candidate pending review.
    with qwen_output.open("w", encoding="utf-8", newline="\n") as f:
        for record in candidates:
            interval = record["image_depth_interval_m"]
            target = {
                "description_segments": [
                    {
                        "top_m": segment["top_m"],
                        "base_m": segment["base_m"],
                        "description": segment["source_description"],
                    }
                    for segment in record["source_description_segments"]
                ]
            }
            conversation = {
                "image": record["image"]["local_copy"],
                "conversations": [
                    {
                        "from": "human",
                        "value": (
                            "<image>\nОпиши керн на фотографии в пределах интервала "
                            f"{interval['top']}–{interval['base']} м. "
                            "Верни только JSON с ключом description_segments; для каждого "
                            "слоя укажи top_m, base_m и description. Не назначай среду "
                            "осадконакопления и не добавляй признаки, которых нет в исходном описании."
                        ),
                    },
                    {"from": "gpt", "value": json.dumps(target, ensure_ascii=False)},
                ],
                "metadata": {
                    "record_id": record["id"],
                    "review_status": record["review_status"],
                    "training_approved": record["training_approved"],
                    "environment_label": record["environment_label"],
                    "description_coverage_fraction": record["description_coverage_fraction"],
                },
            }
            f.write(json.dumps(conversation, ensure_ascii=False) + "\n")

    summary = {
        "source_well": "Уренгойское 41",
        "source_archive": str(ARCHIVE),
        "source_workbook": str(WORKBOOK),
        "daylight_core_photos_scanned": daylight_count,
        "described_layer_rows": len(layers),
        "candidate_pairs": len(candidates),
        "extracted_candidate_image_bytes": extracted_bytes,
        "qwen_format_manifest": str(qwen_output),
        "omitted_photos": len(omitted),
        "candidate_pairs_by_type": dict(Counter(x["pair_type"] for x in candidates)),
        "candidate_pairs_by_formation": dict(Counter(f for x in candidates for f in x["formation_candidates"])),
        "minimum_description_coverage": MIN_COVERAGE,
        "note": "Candidates are interval-linked source pairs, not geologist-approved VLM training labels. No depositional environment was inferred or assigned.",
    }
    (PROJECT_ROOT / "data" / "pairing_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (PROJECT_ROOT / "data" / "omitted_photos.jsonl").write_text(
        "".join(json.dumps(x, ensure_ascii=False) + "\n" for x in omitted), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
